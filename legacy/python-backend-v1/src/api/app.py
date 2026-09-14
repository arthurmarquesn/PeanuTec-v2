from collections import Counter
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Literal
from uuid import uuid4
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from fastapi import FastAPI, Header, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field, StrictBool, field_validator, model_validator
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from src.data_sources.open_meteo_client import fetch_weather_context
from src.engine import risk_engine
from src.repositories import (
    analysis_history_repository,
    calendar_events_repository,
    fields_repository,
    inspections_repository,
    products_repository,
    spray_applications_repository,
    users_repository,
)
from src.services import auth_service, defense_estimator, geocoding_service


SUPPORTED_DISEASES = [
    {
        "id": "Mancha-preta",
        "name": "Mancha-preta do amendoim",
        "pathogen": "Cercosporidium personatum / Nothopassalora personata",
        "status": "validated",
    },
    {
        "id": "Mancha-castanha",
        "name": "Mancha-castanha do amendoim",
        "pathogen": "Cercospora arachidicola",
        "status": "validated",
    },
]

ALLOWED_ORIGINS = [
    "http://localhost:3000",
    "http://127.0.0.1:3000",
]

SUPPORTED_DISEASE_IDS = ["Mancha-preta", "Mancha-castanha"]
SUPPORTED_INSPECTION_DISEASES = [
    "Mancha-preta do amendoim",
    "Mancha-castanha do amendoim",
    "Mancha-preta",
    "Mancha-castanha",
]
SUPPORTED_CROP_STATUSES = ["em_campo", "pre_arranquio", "arrancado", "colhido"]
SUPPORTED_CROP_STAGES = [
    "plantio",
    "emergencia_estabelecimento",
    "vegetativo",
    "florescimento",
    "enchimento_vagens",
    "pre_arranquio",
    "arranquio",
    "colheita",
    "pos_colheita",
]
SUPPORTED_VISUAL_SEVERITIES = ["nenhuma", "baixa", "media", "alta"]
SUPPORTED_DEFOLIATION_LEVELS = ["nenhuma", "baixa", "media", "alta"]
SUPPORTED_DISEASE_INCIDENCE_LEVELS = ["nenhuma", "baixa", "media", "alta"]
SUPPORTED_HISTORICAL_PRESSURES = ["baixa", "media", "alta"]
SUPPORTED_INSPECTION_GENERAL_STATUSES = ["boa", "regular", "atencao", "critica"]
SUPPORTED_PROBLEM_DISTRIBUTIONS = [
    "ausente",
    "localizado",
    "reboleiras",
    "espalhado",
    "generalizado",
]
SUPPORTED_WEED_PRESSURES = ["nenhuma", "baixa", "media", "alta"]
SUPPORTED_SOIL_CONDITIONS = [
    "seco",
    "adequado",
    "umido",
    "encharcado",
    "compactado",
    "nao_avaliado",
]
SUPPORTED_ACTIONS_TAKEN = [
    "nenhuma",
    "monitorar",
    "consultar_responsavel",
    "manejo_realizado",
]
SUPPORTED_CALENDAR_EVENT_TYPES = [
    "pulverizacao",
    "inspecao",
    "monitoramento",
    "reaplicacao_prevista",
    "observacao",
]
SUPPORTED_PRODUCT_TYPES = [
    "fungicida",
    "inseticida",
    "acaricida",
    "herbicida",
    "outro",
]
CALENDAR_EVENT_COLOR_KEYS = {
    "pulverizacao": "yellow",
    "inspecao": "green",
    "monitoramento": "blue",
    "reaplicacao_prevista": "purple",
    "observacao": "gray",
}
CALENDAR_EVENT_INITIAL_STATUSES = {
    "pulverizacao": "realizado",
    "inspecao": "realizado",
    "monitoramento": "previsto",
    "reaplicacao_prevista": "previsto",
    "observacao": "informativo",
}
CALENDAR_EXPORT_HEADERS = [
    "Data",
    "Data final",
    "Talh\u00e3o",
    "Tipo de evento",
    "Status",
    "Produto",
    "Tipo do produto",
    "Alvo",
    "Dose",
    "Respons\u00e1vel",
    "Intervalo planejado",
    "Reaplica\u00e7\u00e3o prevista",
    "Observa\u00e7\u00f5es",
]
CalendarEventTypeFilter = Literal[
    "pulverizacao",
    "inspecao",
    "monitoramento",
    "reaplicacao_prevista",
    "observacao",
]
ProductTypeFilter = Literal[
    "fungicida",
    "inseticida",
    "acaricida",
    "herbicida",
    "outro",
]
TIMEZONE = "America/Sao_Paulo"
TECHNICAL_REPORT_SAFETY_NOTE = (
    "Este relatório organiza dados operacionais do talhão e apoia a tomada de "
    "decisão. Ele não representa diagnóstico definitivo de doença, não recomenda "
    "defensivo e não substitui a avaliação técnica responsável."
)
OPERATIONAL_CONTEXT_SAFETY_NOTE = (
    "Esta leitura organiza dados operacionais do talhão e apoia o acompanhamento "
    "da propriedade. Ela não representa diagnóstico definitivo, laudo técnico ou "
    "recomendação automática de manejo."
)
RISK_CLASS_PRIORITY = {
    "CRÍTICO": 4,
    "CRITICO": 4,
    "CRÃTICO": 4,
    "ALTO": 3,
    "MODERADO": 2,
    "BAIXO": 1,
}


class FieldAnalysisRequest(BaseModel):
    nome: str
    cidade: str
    latitude: float = Field(ge=-90, le=90)
    longitude: float = Field(ge=-180, le=180)
    data_plantio: date
    cultura: str
    doenca_alvo: str
    status_lavoura: str = "em_campo"

    @field_validator(
        "nome",
        "cidade",
        "cultura",
        "doenca_alvo",
        "status_lavoura",
        mode="before",
    )
    @classmethod
    def strip_string_fields(cls, value):
        if isinstance(value, str):
            return value.strip()
        return value

    @field_validator("nome", "cidade")
    @classmethod
    def reject_empty_strings(cls, value: str) -> str:
        if not value:
            raise ValueError("campo nao pode ser vazio")
        return value

    @field_validator("cultura")
    @classmethod
    def validate_crop(cls, value: str) -> str:
        if value != "Amendoim":
            raise ValueError("cultura deve ser Amendoim")
        return value

    @field_validator("doenca_alvo")
    @classmethod
    def validate_target_disease(cls, value: str) -> str:
        if value not in SUPPORTED_DISEASE_IDS:
            raise ValueError("doenca_alvo deve ser Mancha-preta ou Mancha-castanha")
        return value

    @field_validator("status_lavoura")
    @classmethod
    def validate_crop_status(cls, value: str) -> str:
        if value not in SUPPORTED_CROP_STATUSES:
            raise ValueError(
                "status_lavoura deve ser em_campo, pre_arranquio, arrancado ou colhido"
            )
        return value

    @field_validator("data_plantio")
    @classmethod
    def reject_future_planting_date(cls, value: date) -> date:
        if value > date.today():
            raise ValueError("data_plantio nao pode ser futura")
        return value


class FieldRegistrationRequest(BaseModel):
    nome: str
    cidade: str
    cultura: str
    data_plantio: date
    status_lavoura: str = "em_campo"
    doencas_monitoradas: list[str] = Field(min_length=1)
    previous_crop: str | None = None
    crop_rotation: StrictBool | None = None
    peanut_repetition_years: int | None = None
    had_disease_incidence: StrictBool | None = None
    previous_diseases: str | list[str] | None = None
    disease_incidence_level: Literal["nenhuma", "baixa", "media", "alta"] | None = None
    historical_pressure: Literal["baixa", "media", "alta"] | None = None
    agronomic_history_notes: str | None = None

    @field_validator(
        "nome",
        "cidade",
        "cultura",
        "status_lavoura",
        "previous_crop",
        "disease_incidence_level",
        "historical_pressure",
        "agronomic_history_notes",
        mode="before",
    )
    @classmethod
    def strip_string_fields(cls, value):
        if isinstance(value, str):
            return value.strip()
        return value

    @field_validator(
        "previous_crop",
        "disease_incidence_level",
        "historical_pressure",
        "agronomic_history_notes",
    )
    @classmethod
    def empty_optional_string_to_none(cls, value: str | None) -> str | None:
        if value == "":
            return None
        return value

    @field_validator("doencas_monitoradas", mode="before")
    @classmethod
    def strip_monitored_diseases(cls, value):
        if isinstance(value, list):
            return [
                disease.strip() if isinstance(disease, str) else disease
                for disease in value
            ]
        return value

    @field_validator("previous_diseases", mode="before")
    @classmethod
    def normalize_previous_diseases(cls, value):
        if value == "":
            return None

        if isinstance(value, list):
            return [
                disease.strip() if isinstance(disease, str) else disease
                for disease in value
                if not isinstance(disease, str) or disease.strip()
            ]

        if isinstance(value, str):
            return value.strip()

        return value

    @field_validator("nome", "cidade")
    @classmethod
    def reject_empty_strings(cls, value: str) -> str:
        if not value:
            raise ValueError("campo nao pode ser vazio")
        return value

    @field_validator("cultura")
    @classmethod
    def validate_crop(cls, value: str) -> str:
        if value != "Amendoim":
            raise ValueError("cultura deve ser Amendoim")
        return value

    @field_validator("status_lavoura")
    @classmethod
    def validate_crop_status(cls, value: str) -> str:
        if value not in SUPPORTED_CROP_STATUSES:
            raise ValueError(
                "status_lavoura deve ser em_campo, pre_arranquio, arrancado ou colhido"
            )
        return value

    @field_validator("doencas_monitoradas")
    @classmethod
    def validate_monitored_diseases(cls, value: list[str]) -> list[str]:
        invalid_diseases = [
            disease for disease in value if disease not in SUPPORTED_DISEASE_IDS
        ]

        if invalid_diseases:
            raise ValueError("doencas_monitoradas contem doenca nao suportada")

        return value

    @field_validator("data_plantio")
    @classmethod
    def reject_future_planting_date(cls, value: date) -> date:
        if value > date.today():
            raise ValueError("data_plantio nao pode ser futura")
        return value

    @field_validator("peanut_repetition_years")
    @classmethod
    def validate_peanut_repetition_years(cls, value: int | None) -> int | None:
        if value is not None and value < 0:
            raise ValueError("peanut_repetition_years nao pode ser negativo")
        return value


class CropStageUpdateRequest(BaseModel):
    stage: Literal[
        "plantio",
        "emergencia_estabelecimento",
        "vegetativo",
        "florescimento",
        "enchimento_vagens",
        "pre_arranquio",
        "arranquio",
        "colheita",
        "pos_colheita",
    ] | None
    notes: str | None = None

    @field_validator("notes", mode="before")
    @classmethod
    def strip_notes(cls, value):
        if isinstance(value, str):
            return value.strip() or None
        return value


class FieldInspectionRequest(BaseModel):
    disease: str
    symptoms_found: StrictBool
    visual_severity: str
    defoliation_level: str
    action_taken: str
    notes: str = ""
    inspected_at: datetime | None = None
    general_status: str | None = None
    problem_distribution: str | None = None
    pests_found: StrictBool | None = None
    pest_notes: str | None = None
    weeds_found: StrictBool | None = None
    weed_pressure: str | None = None
    soil_condition: str | None = None
    return_needed: StrictBool | None = None
    return_days: int | None = None
    observed_area: str | None = None
    responsible: str | None = None

    @field_validator(
        "disease",
        "visual_severity",
        "defoliation_level",
        "action_taken",
        "notes",
        "general_status",
        "problem_distribution",
        "pest_notes",
        "weed_pressure",
        "soil_condition",
        "observed_area",
        "responsible",
        mode="before",
    )
    @classmethod
    def strip_string_fields(cls, value):
        if isinstance(value, str):
            return value.strip()
        return value

    @field_validator("disease")
    @classmethod
    def validate_disease(cls, value: str) -> str:
        if not value:
            raise ValueError("disease nao pode ser vazio")

        if value not in SUPPORTED_INSPECTION_DISEASES:
            raise ValueError("disease deve ser uma doenca suportada")

        return value

    @field_validator("visual_severity")
    @classmethod
    def validate_visual_severity(cls, value: str) -> str:
        if value not in SUPPORTED_VISUAL_SEVERITIES:
            raise ValueError("visual_severity invalida")

        return value

    @field_validator("defoliation_level")
    @classmethod
    def validate_defoliation_level(cls, value: str) -> str:
        if value not in SUPPORTED_DEFOLIATION_LEVELS:
            raise ValueError("defoliation_level invalido")

        return value

    @field_validator("action_taken")
    @classmethod
    def validate_action_taken(cls, value: str) -> str:
        if value not in SUPPORTED_ACTIONS_TAKEN:
            raise ValueError("action_taken invalido")

        return value

    @field_validator(
        "general_status",
        "problem_distribution",
        "pest_notes",
        "weed_pressure",
        "soil_condition",
        "observed_area",
        "responsible",
        mode="before",
    )
    @classmethod
    def empty_optional_string_to_none(cls, value):
        if value == "":
            return None
        return value

    @field_validator("general_status")
    @classmethod
    def validate_general_status(cls, value: str | None) -> str | None:
        if value is not None and value not in SUPPORTED_INSPECTION_GENERAL_STATUSES:
            raise ValueError("general_status invalido")
        return value

    @field_validator("problem_distribution")
    @classmethod
    def validate_problem_distribution(cls, value: str | None) -> str | None:
        if value is not None and value not in SUPPORTED_PROBLEM_DISTRIBUTIONS:
            raise ValueError("problem_distribution invalida")
        return value

    @field_validator("weed_pressure")
    @classmethod
    def validate_weed_pressure(cls, value: str | None) -> str | None:
        if value is not None and value not in SUPPORTED_WEED_PRESSURES:
            raise ValueError("weed_pressure invalida")
        return value

    @field_validator("soil_condition")
    @classmethod
    def validate_soil_condition(cls, value: str | None) -> str | None:
        if value is not None and value not in SUPPORTED_SOIL_CONDITIONS:
            raise ValueError("soil_condition invalida")
        return value

    @field_validator("return_days")
    @classmethod
    def validate_return_days(cls, value: int | None) -> int | None:
        if value is not None and value < 0:
            raise ValueError("return_days nao pode ser negativo")
        return value


class InspectionScopeRequest(FieldInspectionRequest):
    scope: Literal["selected", "all"]
    field_ids: list[str] | None = None

    @field_validator("field_ids", mode="before")
    @classmethod
    def strip_field_ids(cls, value):
        if isinstance(value, list):
            return [
                field_id.strip() if isinstance(field_id, str) else field_id
                for field_id in value
            ]

        return value


class SprayApplicationRequest(BaseModel):
    application_date: datetime | None = None
    product_id: str | None = None
    product: str
    product_type: Literal[
        "fungicida",
        "inseticida",
        "acaricida",
        "herbicida",
        "outro",
    ] | None = None
    target: str
    dose: str
    responsible: str = ""
    planned_interval_days: int = 12
    notes: str = ""
    generate_reapplication: StrictBool = False
    reapplication_interval_days: int | None = None
    reapplication_date: date | None = None
    reapplication_notes: str | None = None

    @model_validator(mode="before")
    @classmethod
    def validate_reapplication_configuration(cls, data):
        if not isinstance(data, dict):
            return data

        generate_reapplication = data.get("generate_reapplication") is True

        if generate_reapplication:
            has_reapplication_date = bool(data.get("reapplication_date"))
            has_reapplication_interval = data.get("reapplication_interval_days") not in (
                None,
                "",
            )
            has_planned_interval = data.get("planned_interval_days") not in (None, "")

            if (
                not has_reapplication_date
                and not has_reapplication_interval
                and not has_planned_interval
            ):
                raise ValueError(
                    "Informe a data prevista ou um intervalo para gerar a reaplicacao."
                )

        return data

    @field_validator(
        "product",
        "product_id",
        "product_type",
        "target",
        "dose",
        "responsible",
        "notes",
        "reapplication_notes",
        mode="before",
    )
    @classmethod
    def strip_string_fields(cls, value):
        if isinstance(value, str):
            return value.strip()
        return value

    @field_validator("product_id", "product_type")
    @classmethod
    def empty_optional_string_to_none(cls, value: str | None) -> str | None:
        if value == "":
            return None
        return value

    @field_validator("reapplication_interval_days", mode="before")
    @classmethod
    def empty_reapplication_interval_to_none(cls, value):
        if value == "":
            return None
        return value

    @field_validator("reapplication_date", mode="before")
    @classmethod
    def empty_reapplication_date_to_none(cls, value):
        if value == "":
            return None
        return value

    @field_validator("product", "target", "dose")
    @classmethod
    def reject_empty_required_strings(cls, value: str) -> str:
        if not value:
            raise ValueError("campo nao pode ser vazio")
        return value

    @field_validator("planned_interval_days")
    @classmethod
    def validate_planned_interval_days(cls, value: int) -> int:
        if value <= 0:
            raise ValueError("planned_interval_days deve ser inteiro positivo")
        return value

    @field_validator("reapplication_interval_days")
    @classmethod
    def validate_reapplication_interval_days(cls, value: int | None) -> int | None:
        if value is not None and value <= 0:
            raise ValueError("reapplication_interval_days deve ser inteiro positivo")
        return value

    @field_validator("application_date")
    @classmethod
    def normalize_application_date(cls, value: datetime | None) -> datetime | None:
        if value is None:
            return value

        normalized = normalize_datetime_to_local_timezone(value)

        if normalized > datetime.now(ZoneInfo(TIMEZONE)):
            raise ValueError("application_date nao pode ser futura")

        return normalized


class CalendarEventRequest(BaseModel):
    event_type: Literal[
        "pulverizacao",
        "inspecao",
        "monitoramento",
        "reaplicacao_prevista",
        "observacao",
    ]
    title: str
    field_id: str | None = None
    date: date
    end_date: date | None = None
    product: str = ""
    product_type: Literal[
        "fungicida",
        "inseticida",
        "acaricida",
        "herbicida",
        "outro",
    ] | None = None
    target: str = ""
    planned_interval_days: int | None = None
    notes: str = ""

    @field_validator(
        "title",
        "field_id",
        "product",
        "product_type",
        "target",
        "notes",
        mode="before",
    )
    @classmethod
    def strip_string_fields(cls, value):
        if isinstance(value, str):
            return value.strip()
        return value

    @field_validator("field_id")
    @classmethod
    def empty_field_id_to_none(cls, value: str | None) -> str | None:
        if value == "":
            return None
        return value

    @field_validator("title")
    @classmethod
    def reject_empty_title(cls, value: str) -> str:
        if not value:
            raise ValueError("title nao pode ser vazio")
        return value

    @field_validator("planned_interval_days")
    @classmethod
    def validate_planned_interval_days(cls, value: int | None) -> int | None:
        if value is not None and value <= 0:
            raise ValueError("planned_interval_days deve ser inteiro positivo")
        return value

    @model_validator(mode="after")
    def validate_end_date(self):
        if self.end_date is not None and self.end_date < self.date:
            raise ValueError("end_date nao pode ser anterior a date")
        return self


class ProductRequest(BaseModel):
    name: str
    product_type: Literal[
        "fungicida",
        "inseticida",
        "acaricida",
        "herbicida",
        "outro",
    ]
    active_ingredient: str | None = None
    main_target: str | None = None
    default_defense_days: int | None = None
    notes: str | None = None
    is_active: StrictBool = True

    @field_validator(
        "name",
        "product_type",
        "active_ingredient",
        "main_target",
        "notes",
        mode="before",
    )
    @classmethod
    def strip_string_fields(cls, value):
        if isinstance(value, str):
            return value.strip()
        return value

    @field_validator("active_ingredient", "main_target", "notes")
    @classmethod
    def empty_optional_string_to_none(cls, value: str | None) -> str | None:
        if value == "":
            return None
        return value

    @field_validator("default_defense_days", mode="before")
    @classmethod
    def empty_default_defense_days_to_none(cls, value):
        if value == "":
            return None
        return value

    @field_validator("name")
    @classmethod
    def reject_empty_name(cls, value: str) -> str:
        if not value:
            raise ValueError("name nao pode ser vazio")
        return value

    @field_validator("default_defense_days")
    @classmethod
    def validate_default_defense_days(cls, value: int | None) -> int | None:
        if value is not None and value <= 0:
            raise ValueError("default_defense_days deve ser maior que 0")
        return value


class ProductUpdateRequest(ProductRequest):
    pass


class ProductResponse(BaseModel):
    id: str
    name: str
    product_type: str
    active_ingredient: str | None = None
    main_target: str | None = None
    default_defense_days: int | None = None
    notes: str | None = None
    is_active: bool
    created_at: str | None = None
    updated_at: str | None = None


class RegisterRequest(BaseModel):
    name: str
    email: str
    password: str = Field(min_length=6)

    @model_validator(mode="before")
    @classmethod
    def accept_frontend_field_names(cls, data):
        if not isinstance(data, dict):
            return data

        normalized = dict(data)

        if "name" not in normalized and "nome" in normalized:
            normalized["name"] = normalized["nome"]

        if "password" not in normalized and "senha" in normalized:
            normalized["password"] = normalized["senha"]

        return normalized

    @field_validator("name", "email", "password", mode="before")
    @classmethod
    def strip_string_fields(cls, value):
        if isinstance(value, str):
            return value.strip()
        return value

    @field_validator("name", "email")
    @classmethod
    def reject_empty_strings(cls, value: str) -> str:
        if not value:
            raise ValueError("campo nao pode ser vazio")
        return value


class LoginRequest(BaseModel):
    email: str
    password: str

    @model_validator(mode="before")
    @classmethod
    def accept_frontend_field_names(cls, data):
        if not isinstance(data, dict):
            return data

        normalized = dict(data)

        if "password" not in normalized and "senha" in normalized:
            normalized["password"] = normalized["senha"]

        return normalized

    @field_validator("email", "password", mode="before")
    @classmethod
    def strip_string_fields(cls, value):
        if isinstance(value, str):
            return value.strip()
        return value


def extract_bearer_token(authorization: str | None) -> str:
    if not authorization:
        raise HTTPException(status_code=401, detail="Token ausente")

    scheme, _, token = authorization.partition(" ")

    if scheme.lower() != "bearer" or not token:
        raise HTTPException(status_code=401, detail="Token invalido")

    return token


app = FastAPI(
    title="Peanut Disease Risk Engine",
    version="0.2",
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["Content-Type", "Authorization"],
)


@app.get("/health")
def health_check() -> dict:
    return {
        "status": "ok",
        "service": "peanut-disease-risk-engine",
    }


@app.get("/")
def get_service_metadata() -> dict:
    return {
        "service": "Peanut Disease Risk Engine",
        "version": "0.2",
        "supported_diseases": SUPPORTED_DISEASE_IDS,
    }


@app.get("/doencas")
def list_supported_diseases() -> dict:
    return {
        "supported_diseases": SUPPORTED_DISEASES,
    }


@app.post("/auth/register")
def register_user(payload: RegisterRequest) -> dict:
    try:
        user = auth_service.register_user(
            name=payload.name,
            email=payload.email,
            password=payload.password,
        )
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error

    return auth_service.public_user(user)


@app.post("/auth/login")
def login_user(payload: LoginRequest) -> dict:
    user = auth_service.authenticate_user(
        email=payload.email,
        password=payload.password,
    )

    if user is None:
        raise HTTPException(status_code=401, detail="Credenciais invalidas")

    return {
        "access_token": auth_service.create_access_token(user["id"]),
        "token_type": "bearer",
        "user": auth_service.public_user(user),
    }


@app.get("/auth/me")
def get_current_user(authorization: str | None = Header(default=None)) -> dict:
    token = extract_bearer_token(authorization)
    user = auth_service.get_user_from_token(token)

    if user is None:
        raise HTTPException(status_code=401, detail="Token invalido")

    return auth_service.public_user(user)


def build_product_record(
    product_data: dict,
    product_id: str | None = None,
    current_product: dict | None = None,
) -> dict:
    current_time = datetime.now(ZoneInfo(TIMEZONE)).isoformat(timespec="seconds")

    return {
        "id": product_id or str(uuid4()),
        "name": product_data["name"],
        "product_type": product_data["product_type"],
        "active_ingredient": product_data.get("active_ingredient"),
        "main_target": product_data.get("main_target"),
        "default_defense_days": product_data.get("default_defense_days"),
        "notes": product_data.get("notes"),
        "is_active": product_data.get("is_active", True),
        "created_at": (
            current_product.get("created_at")
            if current_product is not None
            else current_time
        ),
        "updated_at": current_time,
    }


@app.get("/produtos", response_model=list[ProductResponse])
def list_products(active_only: bool = False) -> list[dict]:
    return products_repository.list_products(active_only=active_only)


@app.post("/produtos", response_model=ProductResponse)
def create_product(product: ProductRequest) -> dict:
    product_record = build_product_record(product.model_dump(mode="json"))
    return products_repository.create_product(product_record)


@app.get("/produtos/{product_id}", response_model=ProductResponse)
def get_product(product_id: str) -> dict:
    product = products_repository.get_product(product_id)

    if product is None:
        raise HTTPException(status_code=404, detail="Produto nao encontrado")

    return product


@app.put("/produtos/{product_id}", response_model=ProductResponse)
def update_product(product_id: str, product: ProductUpdateRequest) -> dict:
    current_product = products_repository.get_product(product_id)

    if current_product is None:
        raise HTTPException(status_code=404, detail="Produto nao encontrado")

    product_record = build_product_record(
        product.model_dump(mode="json"),
        product_id=product_id,
        current_product=current_product,
    )
    updated_product = products_repository.update_product(product_id, product_record)

    if updated_product is None:
        raise HTTPException(status_code=404, detail="Produto nao encontrado")

    return updated_product


@app.delete("/produtos/{product_id}", response_model=ProductResponse)
def deactivate_product(product_id: str) -> dict:
    product = products_repository.deactivate_product(product_id)

    if product is None:
        raise HTTPException(status_code=404, detail="Produto nao encontrado")

    return product


def build_field_record(field_data: dict, field_id: str | None = None) -> dict:
    coordinates = geocoding_service.fetch_city_coordinates(field_data["cidade"])

    return {
        "id": field_id or str(uuid4()),
        "nome": field_data["nome"],
        "cidade": field_data["cidade"],
        "latitude": coordinates["latitude"],
        "longitude": coordinates["longitude"],
        "cultura": field_data["cultura"],
        "data_plantio": field_data["data_plantio"],
        "status_lavoura": field_data["status_lavoura"],
        "doencas_monitoradas": field_data["doencas_monitoradas"],
        "previous_crop": field_data.get("previous_crop"),
        "crop_rotation": field_data.get("crop_rotation"),
        "peanut_repetition_years": field_data.get("peanut_repetition_years"),
        "had_disease_incidence": field_data.get("had_disease_incidence"),
        "previous_diseases": field_data.get("previous_diseases"),
        "disease_incidence_level": field_data.get("disease_incidence_level"),
        "historical_pressure": field_data.get("historical_pressure"),
        "agronomic_history_notes": field_data.get("agronomic_history_notes"),
        "manual_crop_stage": field_data.get("manual_crop_stage"),
        "crop_stage_updated_at": field_data.get("crop_stage_updated_at"),
        "crop_stage_notes": field_data.get("crop_stage_notes"),
    }


@app.get("/talhoes")
def list_registered_fields() -> list[dict]:
    return fields_repository.list_fields()


@app.post("/talhoes")
def create_registered_field(field: FieldRegistrationRequest) -> dict:
    field_data = field.model_dump(mode="json")

    try:
        field_record = build_field_record(field_data)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error

    return fields_repository.create_field(field_record)


@app.get("/talhoes/{field_id}")
def get_registered_field(field_id: str) -> dict:
    field = fields_repository.get_field(field_id)

    if field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    return field


@app.put("/talhoes/{field_id}")
def update_registered_field(
    field_id: str,
    field: FieldRegistrationRequest,
) -> dict:
    current_field = fields_repository.get_field(field_id)

    if current_field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    field_data = field.model_dump(mode="json")
    field_data.update(
        {
            "manual_crop_stage": current_field.get("manual_crop_stage"),
            "crop_stage_updated_at": current_field.get("crop_stage_updated_at"),
            "crop_stage_notes": current_field.get("crop_stage_notes"),
        }
    )

    try:
        field_record = build_field_record(field_data, field_id=field_id)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error

    updated_field = fields_repository.update_field(field_id, field_record)

    if updated_field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    return updated_field


@app.patch("/talhoes/{field_id}/fase-lavoura")
def update_field_crop_stage(
    field_id: str,
    crop_stage: CropStageUpdateRequest,
) -> dict:
    field = fields_repository.get_field(field_id)

    if field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    updated_field = {
        **field,
        "manual_crop_stage": crop_stage.stage,
        "crop_stage_notes": crop_stage.notes if crop_stage.stage else None,
        "crop_stage_updated_at": (
            datetime.now(ZoneInfo(TIMEZONE)).isoformat(timespec="seconds")
            if crop_stage.stage
            else None
        ),
    }
    saved_field = fields_repository.update_field(field_id, updated_field)

    if saved_field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    return {
        "field_id": field_id,
        "manual_crop_stage": saved_field.get("manual_crop_stage"),
        "crop_stage_updated_at": saved_field.get("crop_stage_updated_at"),
        "crop_stage_notes": saved_field.get("crop_stage_notes"),
        "crop_stage_context": build_crop_stage_context(saved_field),
    }


@app.delete("/talhoes/{field_id}")
def delete_registered_field(field_id: str) -> dict:
    deleted = fields_repository.delete_field(field_id)

    if not deleted:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    return {
        "message": "Talhao removido com sucesso",
        "id": field_id,
    }


def build_analysis_payload_from_field(field: dict, disease_id: str) -> dict:
    return {
        "nome": field["nome"],
        "cidade": field["cidade"],
        "latitude": field["latitude"],
        "longitude": field["longitude"],
        "data_plantio": field["data_plantio"],
        "cultura": field["cultura"],
        "doenca_alvo": disease_id,
        "status_lavoura": field["status_lavoura"],
    }


def run_compact_analyses_for_field(field: dict) -> list[dict]:
    analyses = []

    for disease_id in field.get("doencas_monitoradas", []):
        analysis_payload = build_analysis_payload_from_field(field, disease_id)
        analyses.append(risk_engine.run_disease_analysis_compact(analysis_payload))

    return sorted(
        analyses,
        key=lambda analysis: analysis["risk"]["agronomic_index"],
        reverse=True,
    )


@app.post("/talhoes/{field_id}/analisar")
def analyze_registered_field(field_id: str) -> dict:
    field = fields_repository.get_field(field_id)

    if field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    try:
        analyses = run_compact_analyses_for_field(field)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=f"Erro ao executar analise do talhao: {error}",
        ) from error

    for analysis in analyses:
        analysis_history_repository.create_analysis_history_record(
            field_id=field["id"],
            field_name=field["nome"],
            compact_analysis=analysis,
        )

    return {
        "field_id": field["id"],
        "field_name": field["nome"],
        "city": field["cidade"],
        "total_analyses": len(analyses),
        "analyses": analyses,
    }


@app.get("/talhoes/{field_id}/historico-analises")
def list_field_analysis_history(field_id: str) -> dict:
    field = fields_repository.get_field(field_id)

    if field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    history = analysis_history_repository.list_analysis_history_by_field(field_id)
    history.sort(key=lambda record: record["generated_at"], reverse=True)

    return {
        "field_id": field["id"],
        "field_name": field["nome"],
        "total": len(history),
        "history": history,
    }


def build_calendar_event_record(
    *,
    event_type: str,
    title: str,
    event_date: date,
    current_time: datetime,
    end_date: date | None = None,
    field: dict | None = None,
    field_id: str | None = None,
    product: str = "",
    product_type: str | None = None,
    target: str = "",
    planned_interval_days: int | None = None,
    notes: str = "",
    source_type: str | None = None,
    source_id: str | None = None,
) -> dict:
    resolved_field_id = field["id"] if field is not None else field_id

    return {
        "id": str(uuid4()),
        "event_type": event_type,
        "title": title,
        "field_id": resolved_field_id,
        "field_name": field["nome"] if field is not None else None,
        "date": event_date.isoformat(),
        "end_date": (end_date or event_date).isoformat(),
        "product": product,
        "product_type": product_type,
        "target": target,
        "planned_interval_days": planned_interval_days,
        "notes": notes,
        "color_key": CALENDAR_EVENT_COLOR_KEYS[event_type],
        "status": CALENDAR_EVENT_INITIAL_STATUSES[event_type],
        "source_type": source_type,
        "source_id": source_id,
        "created_at": current_time.isoformat(timespec="seconds"),
    }


def validate_calendar_event_field(event: CalendarEventRequest) -> dict | None:
    if event.field_id is None:
        if event.event_type != "observacao":
            raise HTTPException(
                status_code=400,
                detail="field_id e obrigatorio para este tipo de evento",
            )

        return None

    field = fields_repository.get_field(event.field_id)

    if field is None:
        raise HTTPException(
            status_code=400,
            detail=f"Talhao nao encontrado: {event.field_id}",
        )

    return field


def parse_calendar_event_date(event: dict) -> date:
    return date.fromisoformat(event["date"])


def filter_calendar_events(
    events: list[dict],
    start_date: date | None = None,
    end_date: date | None = None,
    field_id: str | None = None,
    event_type: CalendarEventTypeFilter | None = None,
    product_type: ProductTypeFilter | None = None,
) -> list[dict]:
    filtered_events = list(events)

    if start_date is not None:
        filtered_events = [
            event
            for event in filtered_events
            if parse_calendar_event_date(event) >= start_date
        ]

    if end_date is not None:
        filtered_events = [
            event
            for event in filtered_events
            if parse_calendar_event_date(event) <= end_date
        ]

    if field_id is not None:
        filtered_events = [
            event for event in filtered_events if event.get("field_id") == field_id
        ]

    if event_type is not None:
        filtered_events = [
            event for event in filtered_events if event.get("event_type") == event_type
        ]

    if product_type is not None:
        filtered_events = [
            event
            for event in filtered_events
            if event.get("product_type") == product_type
        ]

    filtered_events.sort(
        key=lambda event: (
            event.get("date", ""),
            event.get("end_date") or event.get("date", ""),
            event.get("title", ""),
        )
    )

    return filtered_events


def calendar_export_value(value) -> str | int:
    if value is None:
        return ""

    return value


def build_calendar_export_workbook(events: list[dict]) -> Workbook:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Calendario de Manejo"
    worksheet.append(CALENDAR_EXPORT_HEADERS)

    header_fill = PatternFill("solid", fgColor="E2E8F0")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for event in events:
        reapplication_date = event.get("reapplication_date") or ""
        if not reapplication_date and event.get("event_type") == "reaplicacao_prevista":
            reapplication_date = event.get("date", "")

        worksheet.append(
            [
                calendar_export_value(event.get("date")),
                calendar_export_value(event.get("end_date") or event.get("date")),
                calendar_export_value(
                    event.get("field_name") or event.get("field_id") or ""
                ),
                calendar_export_value(event.get("event_type")),
                calendar_export_value(event.get("status")),
                calendar_export_value(event.get("product")),
                calendar_export_value(event.get("product_type")),
                calendar_export_value(event.get("target")),
                calendar_export_value(event.get("dose")),
                calendar_export_value(event.get("responsible")),
                calendar_export_value(event.get("planned_interval_days")),
                calendar_export_value(reapplication_date),
                calendar_export_value(event.get("notes")),
            ]
        )

    worksheet.freeze_panes = "A2"
    for column_index, header in enumerate(CALENDAR_EXPORT_HEADERS, start=1):
        max_length = len(header)
        for cell in worksheet[get_column_letter(column_index)]:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max_length + 2,
            36,
        )

    return workbook


@app.post("/calendario/eventos")
def create_calendar_event(event: CalendarEventRequest) -> dict:
    field = validate_calendar_event_field(event)
    current_time = datetime.now(ZoneInfo(TIMEZONE))
    event_record = build_calendar_event_record(
        event_type=event.event_type,
        title=event.title,
        event_date=event.date,
        end_date=event.end_date,
        field=field,
        field_id=event.field_id,
        product=event.product,
        product_type=event.product_type,
        target=event.target,
        planned_interval_days=event.planned_interval_days,
        notes=event.notes,
        current_time=current_time,
    )

    return calendar_events_repository.create_calendar_event(event_record)


@app.get("/calendario/eventos")
def list_calendar_events(
    start_date: date | None = None,
    end_date: date | None = None,
    field_id: str | None = None,
    event_type: CalendarEventTypeFilter | None = None,
    product_type: ProductTypeFilter | None = None,
) -> dict:
    events = filter_calendar_events(
        calendar_events_repository.list_calendar_events(),
        start_date=start_date,
        end_date=end_date,
        field_id=field_id,
        event_type=event_type,
        product_type=product_type,
    )

    return {
        "total": len(events),
        "events": events,
    }


@app.get("/calendario/exportar-excel")
def export_calendar_events_excel(
    start_date: date | None = None,
    end_date: date | None = None,
    field_id: str | None = None,
    event_type: CalendarEventTypeFilter | None = None,
    product_type: ProductTypeFilter | None = None,
) -> StreamingResponse:
    events = filter_calendar_events(
        calendar_events_repository.list_calendar_events(),
        start_date=start_date,
        end_date=end_date,
        field_id=field_id,
        event_type=event_type,
        product_type=product_type,
    )
    workbook = build_calendar_export_workbook(events)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                'attachment; filename="calendario-manejo-peanutec.xlsx"'
            )
        },
    )


@app.get("/calendario/resumo")
def get_calendar_summary() -> dict:
    fields = fields_repository.list_fields()
    events = calendar_events_repository.list_calendar_events()
    current_date = datetime.now(ZoneInfo(TIMEZONE)).date()

    return {
        "active_fields": len(
            [field for field in fields if field["status_lavoura"] == "em_campo"]
        ),
        "events_count": len(events),
        "spray_events_count": len(
            [event for event in events if event["event_type"] == "pulverizacao"]
        ),
        "inspection_events_count": len(
            [event for event in events if event["event_type"] == "inspecao"]
        ),
        "upcoming_attention_count": len(
            [
                event
                for event in events
                if event["event_type"] in ["monitoramento", "reaplicacao_prevista"]
                and event["status"] == "previsto"
                and parse_calendar_event_date(event) >= current_date
            ]
        ),
        "fields_without_recent_inspection": 0,
    }


def create_calendar_event_for_inspection(inspection: dict) -> None:
    if calendar_events_repository.calendar_event_exists(
        "inspection",
        inspection["id"],
        "inspecao",
    ):
        return

    inspected_at = datetime.fromisoformat(inspection["inspected_at"])
    current_time = datetime.now(ZoneInfo(TIMEZONE))
    event_record = build_calendar_event_record(
        event_type="inspecao",
        title=f"Inspecao - {inspection['field_name']}",
        event_date=inspected_at.date(),
        field={
            "id": inspection["field_id"],
            "nome": inspection["field_name"],
        },
        target=inspection["disease"],
        notes=inspection.get("notes", ""),
        source_type="inspection",
        source_id=inspection["id"],
        current_time=current_time,
    )

    calendar_events_repository.create_calendar_event(event_record)


def create_inspection_for_field(
    field: dict,
    inspection: FieldInspectionRequest,
    current_time: datetime | None = None,
) -> dict:
    current_time = current_time or datetime.now(ZoneInfo(TIMEZONE))
    inspection_data = inspection.model_dump(
        mode="json",
        exclude={"scope", "field_ids"},
    )
    inspected_at = inspection.inspected_at or current_time

    inspection_record = {
        "id": str(uuid4()),
        "field_id": field["id"],
        "field_name": field["nome"],
        **inspection_data,
        "inspected_at": inspected_at.isoformat(timespec="seconds"),
        "created_at": current_time.isoformat(timespec="seconds"),
    }

    created_inspection = inspections_repository.create_inspection(inspection_record)
    create_calendar_event_for_inspection(created_inspection)

    return created_inspection



@app.post("/talhoes/{field_id}/inspecoes")
def create_field_inspection(
    field_id: str,
    inspection: FieldInspectionRequest,
) -> dict:
    field = fields_repository.get_field(field_id)

    if field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    return create_inspection_for_field(field, inspection)


@app.post("/inspecoes")
def create_scoped_inspections(inspection: InspectionScopeRequest) -> dict:
    if inspection.scope == "selected":
        if not inspection.field_ids:
            raise HTTPException(
                status_code=400,
                detail="field_ids e obrigatorio quando scope for selected",
            )

        fields = []

        for field_id in inspection.field_ids:
            field = fields_repository.get_field(field_id)

            if field is None:
                raise HTTPException(
                    status_code=400,
                    detail=f"Talhao nao encontrado: {field_id}",
                )

            if field["status_lavoura"] != "em_campo":
                raise HTTPException(
                    status_code=400,
                    detail=f"Talhao nao esta ativo para inspecao: {field_id}",
                )

            fields.append(field)
    else:
        fields = [
            field
            for field in fields_repository.list_fields()
            if field["status_lavoura"] == "em_campo"
        ]

        if not fields:
            raise HTTPException(
                status_code=400,
                detail="Nao ha talhoes ativos para registrar inspecao",
            )

    current_time = datetime.now(ZoneInfo(TIMEZONE))
    created_inspections = [
        create_inspection_for_field(field, inspection, current_time=current_time)
        for field in fields
    ]

    return {
        "scope": inspection.scope,
        "created_count": len(created_inspections),
        "field_ids": [inspection["field_id"] for inspection in created_inspections],
        "inspections": created_inspections,
    }


@app.get("/talhoes/{field_id}/inspecoes")
def list_field_inspections(field_id: str) -> dict:
    field = fields_repository.get_field(field_id)

    if field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    inspections = inspections_repository.list_inspections_by_field(field_id)
    inspections.sort(key=lambda item: item["inspected_at"], reverse=True)

    return {
        "field_id": field["id"],
        "field_name": field["nome"],
        "total": len(inspections),
        "inspections": inspections,
    }


def normalize_datetime_to_local_timezone(value: datetime) -> datetime:
    timezone = ZoneInfo(TIMEZONE)

    if value.tzinfo is None:
        return value.replace(tzinfo=timezone)

    return value.astimezone(timezone)


def calculate_spray_interval(
    application_date: datetime | str,
    planned_interval_days: int | None,
    now: datetime | None = None,
) -> dict:
    if isinstance(application_date, str):
        application_datetime = datetime.fromisoformat(application_date)
    else:
        application_datetime = application_date

    application_datetime = normalize_datetime_to_local_timezone(application_datetime)
    current_time = normalize_datetime_to_local_timezone(
        now or datetime.now(ZoneInfo(TIMEZONE))
    )
    days_since_application = (current_time.date() - application_datetime.date()).days
    planned_interval_days = (
        planned_interval_days or defense_estimator.DEFAULT_PLANNED_INTERVAL_DAYS
    )

    if days_since_application <= planned_interval_days - 3:
        interval_status = "em_dia"
    elif days_since_application <= planned_interval_days:
        interval_status = "atencao"
    else:
        interval_status = "atrasado"

    return {
        "days_since_application": days_since_application,
        "interval_status": interval_status,
    }


def enrich_spray_application_interval(spray_application: dict) -> dict:
    spray_application = {
        "product_id": None,
        "product_type": None,
        "planned_interval_days": None,
        "generate_reapplication": False,
        "reapplication_interval_days": None,
        "reapplication_date": None,
        "reapplication_notes": None,
        **spray_application,
    }
    interval = calculate_spray_interval(
        spray_application["application_date"],
        spray_application.get("planned_interval_days"),
    )

    return {
        **spray_application,
        **interval,
    }


def resolve_spray_reapplication_date(
    spray_application: dict,
    application_datetime: datetime,
) -> date:
    explicit_date = spray_application.get("reapplication_date")

    if explicit_date:
        if isinstance(explicit_date, date):
            return explicit_date

        return date.fromisoformat(explicit_date)

    interval_days = (
        spray_application.get("reapplication_interval_days")
        or spray_application.get("planned_interval_days")
    )

    if not interval_days:
        raise HTTPException(
            status_code=400,
            detail="Informe a data prevista ou um intervalo para gerar a reaplicacao.",
        )

    return application_datetime.date() + timedelta(days=interval_days)


def create_calendar_events_for_spray_application(spray_application: dict) -> None:
    application_datetime = datetime.fromisoformat(spray_application["application_date"])
    current_time = datetime.now(ZoneInfo(TIMEZONE))
    field = {
        "id": spray_application["field_id"],
        "nome": spray_application["field_name"],
    }

    if not calendar_events_repository.calendar_event_exists(
        "spray_application",
        spray_application["id"],
        "pulverizacao",
    ):
        spray_event = build_calendar_event_record(
            event_type="pulverizacao",
            title=f"Pulverizacao - {spray_application['field_name']}",
            event_date=application_datetime.date(),
            field=field,
            product=spray_application["product"],
            product_type=spray_application.get("product_type"),
            target=spray_application["target"],
            planned_interval_days=spray_application["planned_interval_days"],
            notes=spray_application.get("notes", ""),
            source_type="spray_application",
            source_id=spray_application["id"],
            current_time=current_time,
        )
        calendar_events_repository.create_calendar_event(spray_event)

    if not spray_application.get("generate_reapplication", False):
        return

    if not calendar_events_repository.calendar_event_exists(
        "spray_application",
        spray_application["id"],
        "reaplicacao_prevista",
    ):
        reapplication_date = resolve_spray_reapplication_date(
            spray_application,
            application_datetime,
        )
        reapplication_event = build_calendar_event_record(
            event_type="reaplicacao_prevista",
            title=f"Reaplicacao prevista - {spray_application['field_name']}",
            event_date=reapplication_date,
            field=field,
            product=spray_application["product"],
            product_type=spray_application.get("product_type"),
            target=spray_application["target"],
            planned_interval_days=(
                spray_application.get("reapplication_interval_days")
                or spray_application.get("planned_interval_days")
            ),
            notes=(
                spray_application.get("reapplication_notes")
                or "Evento gerado automaticamente a partir da pulverizacao."
            ),
            source_type="spray_application",
            source_id=spray_application["id"],
            current_time=current_time,
        )
        calendar_events_repository.create_calendar_event(reapplication_event)


@app.post("/talhoes/{field_id}/pulverizacoes")
def create_field_spray_application(
    field_id: str,
    spray_application: SprayApplicationRequest,
) -> dict:
    field = fields_repository.get_field(field_id)

    if field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    current_time = datetime.now(ZoneInfo(TIMEZONE))
    spray_application_data = spray_application.model_dump(mode="json")

    if spray_application.product_id is not None:
        selected_product = products_repository.get_product(spray_application.product_id)

        if selected_product is None:
            raise HTTPException(status_code=400, detail="Produto nao encontrado")

        if not selected_product.get("is_active", True):
            raise HTTPException(status_code=400, detail="Produto inativo")

        spray_application_data = {
            **spray_application_data,
            "product": selected_product["name"],
            "product_type": spray_application.product_type
            or selected_product.get("product_type"),
        }

    application_date = spray_application.application_date or current_time
    application_date = normalize_datetime_to_local_timezone(application_date)
    interval = calculate_spray_interval(
        application_date,
        spray_application.planned_interval_days,
        now=current_time,
    )

    spray_application_record = {
        "id": str(uuid4()),
        "field_id": field["id"],
        "field_name": field["nome"],
        **spray_application_data,
        "application_date": application_date.isoformat(timespec="seconds"),
        **interval,
        "created_at": current_time.isoformat(timespec="seconds"),
    }

    created_spray_application = spray_applications_repository.create_spray_application(
        spray_application_record
    )
    create_calendar_events_for_spray_application(created_spray_application)

    return created_spray_application


@app.get("/talhoes/{field_id}/pulverizacoes")
def list_field_spray_applications(field_id: str) -> dict:
    field = fields_repository.get_field(field_id)

    if field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    spray_applications = [
        enrich_spray_application_interval(spray_application)
        for spray_application in spray_applications_repository.list_spray_applications_by_field(
            field_id
        )
    ]
    spray_applications.sort(
        key=lambda item: item["application_date"],
        reverse=True,
    )

    return {
        "field_id": field["id"],
        "field_name": field["nome"],
        "total": len(spray_applications),
        "spray_applications": spray_applications,
    }


def normalize_metric_text(value: str | None) -> str:
    if not value:
        return ""

    return value.strip()


def parse_metric_application_month(spray_application: dict) -> str | None:
    application_date = spray_application.get("application_date")

    if not application_date:
        return None

    try:
        return datetime.fromisoformat(application_date).strftime("%Y-%m")
    except ValueError:
        return None


def build_product_lookup() -> dict[str, dict]:
    return {
        product["id"]: product
        for product in products_repository.list_products()
        if product.get("id")
    }


def resolve_product_metric_identity(
    spray_application: dict,
    products_by_id: dict[str, dict],
) -> dict | None:
    product_id = normalize_metric_text(spray_application.get("product_id"))
    product_name = normalize_metric_text(spray_application.get("product"))
    selected_product = products_by_id.get(product_id) if product_id else None

    if selected_product is not None:
        return {
            "key": f"id:{selected_product['id']}",
            "product_id": selected_product["id"],
            "product": selected_product["name"],
            "product_type": spray_application.get("product_type")
            or selected_product.get("product_type"),
        }

    if not product_name:
        return None

    return {
        "key": f"name:{product_name.lower()}",
        "product_id": product_id or None,
        "product": product_name,
        "product_type": spray_application.get("product_type"),
    }


def increment_counter(counter: Counter, key: str | None) -> None:
    if key:
        counter[key] += 1


def sort_count_rows(rows: list[dict], label_key: str) -> list[dict]:
    return sorted(
        rows,
        key=lambda row: (
            -row["applications_count"],
            str(row.get(label_key) or ""),
        ),
    )


@app.get("/metricas/safra")
def get_season_metrics() -> dict:
    fields = fields_repository.list_fields()
    spray_applications = spray_applications_repository.load_spray_applications()
    products_by_id = build_product_lookup()

    product_counts: Counter = Counter()
    product_metric_by_key: dict[str, dict] = {}
    product_type_counts: Counter = Counter()
    field_counts: Counter = Counter()
    field_name_by_id: dict[str, str] = {}
    target_counts: Counter = Counter()
    month_counts: Counter = Counter()
    valid_intervals = [
        spray_application.get("planned_interval_days")
        for spray_application in spray_applications
        if isinstance(spray_application.get("planned_interval_days"), int)
        and spray_application.get("planned_interval_days") > 0
    ]

    for field in fields:
        if field.get("id"):
            field_name_by_id[field["id"]] = field.get("nome", "")

    for spray_application in spray_applications:
        product_metric = resolve_product_metric_identity(
            spray_application,
            products_by_id,
        )

        if product_metric is not None:
            product_counts[product_metric["key"]] += 1
            product_metric_by_key[product_metric["key"]] = product_metric

        increment_counter(
            product_type_counts,
            normalize_metric_text(
                spray_application.get("product_type")
                or (product_metric.get("product_type") if product_metric else None)
            ),
        )

        field_id = normalize_metric_text(spray_application.get("field_id"))
        field_name = normalize_metric_text(spray_application.get("field_name"))

        if field_id:
            field_counts[field_id] += 1

            if field_name:
                field_name_by_id[field_id] = field_name

        increment_counter(
            target_counts,
            normalize_metric_text(spray_application.get("target")),
        )
        increment_counter(month_counts, parse_metric_application_month(spray_application))

    top_products = sort_count_rows(
        [
            {
                "product_id": product_metric_by_key[key]["product_id"],
                "product": product_metric_by_key[key]["product"],
                "product_type": product_metric_by_key[key]["product_type"],
                "applications_count": count,
            }
            for key, count in product_counts.items()
        ],
        "product",
    )
    product_type_distribution = sort_count_rows(
        [
            {
                "product_type": product_type,
                "applications_count": count,
            }
            for product_type, count in product_type_counts.items()
        ],
        "product_type",
    )
    top_fields_by_applications = sort_count_rows(
        [
            {
                "field_id": field_id,
                "field_name": field_name_by_id.get(field_id, field_id),
                "applications_count": count,
            }
            for field_id, count in field_counts.items()
        ],
        "field_name",
    )
    top_targets = sort_count_rows(
        [
            {
                "target": target,
                "applications_count": count,
            }
            for target, count in target_counts.items()
        ],
        "target",
    )
    spray_applications_by_month = [
        {
            "month": month,
            "applications_count": month_counts[month],
        }
        for month in sorted(month_counts)
    ]
    average_planned_interval_days = (
        round(sum(valid_intervals) / len(valid_intervals), 2)
        if valid_intervals
        else 0
    )

    return {
        "summary": {
            "total_fields": len(fields),
            "active_fields": len(
                [
                    field
                    for field in fields
                    if field.get("status_lavoura") == "em_campo"
                ]
            ),
            "total_spray_applications": len(spray_applications),
            "total_products_used": len(product_counts),
            "average_planned_interval_days": average_planned_interval_days,
        },
        "top_products": top_products,
        "product_type_distribution": product_type_distribution,
        "top_fields_by_applications": top_fields_by_applications,
        "top_targets": top_targets,
        "spray_applications_by_month": spray_applications_by_month,
    }


SITUATION_LABELS = {
    "prioridade_maxima": "Prioridade maxima",
    "alta_atencao": "Alta atencao",
    "monitorar_resposta": "Monitorar resposta",
    "monitorar": "Monitorar",
    "estavel": "Estavel",
    "sem_prioridade_operacional": "Sem prioridade operacional imediata",
}

SITUATION_ACTIONS = {
    "prioridade_maxima": "Verificar o talhão com prioridade e revisar o intervalo de pulverização com o responsável técnico.",
    "alta_atencao": "Realizar inspeção de campo e acompanhar a evolução do risco nos próximos dias.",
    "monitorar_resposta": "Manter monitoramento e registrar inspeção para acompanhar resposta do manejo.",
    "monitorar": "Realizar inspeção de campo e acompanhar a evolução do risco nos próximos dias.",
    "estavel": "Manter acompanhamento de rotina.",
    "sem_prioridade_operacional": "Talhão sem prioridade operacional imediata para manejo foliar.",
}

PRIORITY_BLOCK_LIMITS = {
    "defense": 30,
    "inspection": 25,
    "operational": 20,
    "history": 15,
    "risk": 10,
}
PRIORITY_LABELS = {
    "stable": "Estavel",
    "monitoring": "Monitoramento",
    "high_attention": "Alta atencao",
    "maximum": "Prioridade maxima",
}
PRIORITY_TO_LEGACY_PRIORITY = {
    "Estavel": "BAIXA PRIORIDADE",
    "Monitoramento": "MONITORAR",
    "Alta atencao": "ALTA PRIORIDADE",
    "Prioridade maxima": "PRIORIDADE MAXIMA",
}
PRIORITY_TO_SITUATION = {
    "Estavel": "estavel",
    "Monitoramento": "monitorar",
    "Alta atencao": "alta_atencao",
    "Prioridade maxima": "prioridade_maxima",
}
PRIORITY_RECENT_INSPECTION_WINDOW_DAYS = 7
PRIORITY_RECENT_RECORD_WINDOW_DAYS = 14
CALENDAR_ACTIVE_STATUSES = {"previsto", "informativo"}


def normalize_match_text(value: str | None) -> str:
    if not value:
        return ""

    replacements = {
        "á": "a",
        "à": "a",
        "â": "a",
        "ã": "a",
        "é": "e",
        "ê": "e",
        "í": "i",
        "ó": "o",
        "ô": "o",
        "õ": "o",
        "ú": "u",
        "ç": "c",
    }
    normalized = value.lower()

    for source, target in replacements.items():
        normalized = normalized.replace(source, target)

    normalized = normalized.replace("-", " ").replace("_", " ")

    return " ".join(normalized.split())


def is_critical_or_high_risk(risk_classification: str) -> bool:
    normalized = normalize_match_text(risk_classification).upper()

    return normalized in {"ALTO", "CRITICO", "CRÃTICO", "CRÃƒÂTICO"}


def is_low_risk(risk_classification: str) -> bool:
    return normalize_match_text(risk_classification).upper() == "BAIXO"


def is_moderate_risk(risk_classification: str) -> bool:
    return normalize_match_text(risk_classification).upper() == "MODERADO"


def build_priority_reason(
    *,
    code: str,
    label: str,
    impact: str,
    points: int,
    description: str,
) -> dict:
    return {
        "code": code,
        "label": label,
        "impact": impact,
        "points": points,
        "description": description,
    }


def cap_priority_reasons(reasons: list[dict], limit: int) -> tuple[int, list[dict]]:
    total = sum(reason["points"] for reason in reasons)

    if total <= limit:
        return total, reasons

    excess = total - limit
    capped_reasons = [dict(reason) for reason in reasons]

    for reason in reversed(capped_reasons):
        if excess <= 0:
            break

        reduction = min(reason["points"], excess)
        reason["points"] -= reduction
        excess -= reduction

    capped_reasons = [reason for reason in capped_reasons if reason["points"] > 0]

    return limit, capped_reasons


def priority_label_from_score(score: int) -> str:
    if score >= 75:
        return PRIORITY_LABELS["maximum"]
    if score >= 50:
        return PRIORITY_LABELS["high_attention"]
    if score >= 25:
        return PRIORITY_LABELS["monitoring"]

    return PRIORITY_LABELS["stable"]


def confidence_label_from_score(score: int) -> str:
    if score >= 75:
        return "Alta"
    if score >= 45:
        return "Média"

    return "Baixa"


DISPLAY_VALUE_LABELS = {
    "alta": "Alta",
    "muito_alta": "Muito alta",
    "media": "Média",
    "baixa": "Baixa",
    "nenhuma": "Nenhuma",
    "vencida": "Vencida",
    "sem_registro": "Sem registro",
    "em_dia": "Em dia",
    "atencao": "Atenção",
    "atrasado": "Atrasado",
    "manejo_realizado": "Manejo realizado",
    "consultar_responsavel": "Consultar responsável",
    "monitorar": "Monitorar",
    "critico": "Crítico",
    "alto": "Alto",
    "moderado": "Moderado",
    "baixo": "Baixo",
    "boa": "Boa",
    "regular": "Regular",
    "critica": "Crítica",
    "ausente": "Ausente",
    "localizado": "Localizado",
    "reboleiras": "Reboleiras",
    "espalhado": "Espalhado",
    "generalizado": "Generalizado",
    "seco": "Seco",
    "adequado": "Adequado",
    "umido": "Úmido",
    "encharcado": "Encharcado",
    "compactado": "Compactado",
    "nao_avaliado": "Não avaliado",
}


def format_display_value(value: str | None) -> str:
    if value is None:
        return "Não informado"

    normalized = normalize_match_text(str(value)).replace(" ", "_")

    return DISPLAY_VALUE_LABELS.get(
        normalized,
        str(value).replace("_", " ").capitalize(),
    )


def parse_optional_date(value: str | date | None) -> date | None:
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    try:
        return date.fromisoformat(str(value))
    except ValueError:
        return None


def days_since_optional_datetime(value: str | None, current_date: date) -> int | None:
    parsed_datetime = parse_optional_datetime(value)

    if parsed_datetime is None:
        return None

    return (current_date - parsed_datetime.date()).days


def get_field_calendar_events(field_id: str) -> list[dict]:
    return [
        event
        for event in calendar_events_repository.list_calendar_events()
        if event.get("field_id") == field_id
    ]


def summarize_field_calendar_events(
    events: list[dict],
    current_date: date,
) -> dict:
    overdue_events = []
    overdue_reapplications = []
    today_events = []
    active_alerts = []

    for event in events:
        event_date = parse_optional_date(event.get("date"))

        if event_date is None:
            continue

        status = event.get("status")
        is_active_status = status in CALENDAR_ACTIVE_STATUSES

        if event_date < current_date and is_active_status:
            overdue_events.append(event)

            if event.get("event_type") == "reaplicacao_prevista":
                overdue_reapplications.append(event)

        if event_date == current_date and is_active_status:
            today_events.append(event)

        if is_active_status and event.get("event_type") in {
            "monitoramento",
            "reaplicacao_prevista",
        }:
            active_alerts.append(event)

    return {
        "overdue_events": overdue_events,
        "overdue_reapplications": overdue_reapplications,
        "today_events": today_events,
        "active_alerts": active_alerts,
    }


def calculate_defense_priority(spray_context: dict) -> dict:
    reasons = []
    defense_status = spray_context.get("defense_status")
    interval_status = spray_context.get("interval_status")

    if not spray_context.get("has_spray_record"):
        reasons.append(
            build_priority_reason(
                code="sem_pulverizacao_registrada",
                label="Sem pulverização registrada",
                impact="defesa",
                points=24,
                description="Não há registro de pulverização para estimar a janela operacional do talhão.",
            )
        )
    elif defense_status == "vencida":
        reasons.append(
            build_priority_reason(
                code="defesa_vencida",
                label="Defesa estimada vencida",
                impact="defesa",
                points=20,
                description="A defesa estimada está fora da janela operacional planejada.",
            )
        )
    elif defense_status == "baixa":
        reasons.append(
            build_priority_reason(
                code="defesa_baixa",
                label="Defesa estimada baixa",
                impact="defesa",
                points=14,
                description="A proteção operacional estimada está baixa e exige acompanhamento.",
            )
        )
    elif defense_status == "media":
        reasons.append(
            build_priority_reason(
                code="defesa_media",
                label="Defesa estimada média",
                impact="defesa",
                points=8,
                description="A defesa estimada está em nível intermediário.",
            )
        )

    if interval_status == "atrasado":
        reasons.append(
            build_priority_reason(
                code="intervalo_planejado_vencido",
                label="Intervalo planejado vencido",
                impact="defesa",
                points=10,
                description="O intervalo planejado desde a última aplicação foi ultrapassado.",
            )
        )
    elif interval_status == "atencao":
        reasons.append(
            build_priority_reason(
                code="intervalo_proximo_vencimento",
                label="Intervalo próximo do vencimento",
                impact="defesa",
                points=6,
                description="O intervalo planejado está próximo do limite operacional.",
            )
        )

    points, reasons = cap_priority_reasons(
        reasons,
        PRIORITY_BLOCK_LIMITS["defense"],
    )

    return {
        "points": points,
        "reasons": reasons,
    }


def calculate_inspection_priority(
    inspection_context: dict,
    current_date: date,
) -> dict:
    reasons = []
    days_since_inspection = days_since_optional_datetime(
        inspection_context.get("last_inspection_date")
        or inspection_context.get("inspected_at"),
        current_date,
    )

    if not inspection_context.get("has_inspection_record"):
        reasons.append(
            build_priority_reason(
                code="nunca_inspecionado",
                label="Nunca inspecionado",
                impact="inspecao",
                points=15,
                description="Não há inspeção registrada para apoiar a priorização do talhão.",
            )
        )
    elif (
        days_since_inspection is not None
        and days_since_inspection > PRIORITY_RECENT_INSPECTION_WINDOW_DAYS
    ):
        reasons.append(
            build_priority_reason(
                code="sem_inspecao_recente",
                label="Sem inspeção recente",
                impact="inspecao",
                points=8,
                description="A última inspeção está fora da janela recente de acompanhamento.",
            )
        )

    if inspection_context.get("symptoms_found") is True:
        reasons.append(
            build_priority_reason(
                code="sintomas_presentes",
                label="Presença de sintomas",
                impact="inspecao",
                points=6,
                description="A última inspeção registrou sintomas em campo.",
            )
        )

    general_status = inspection_context.get("general_status")
    general_status_points = {
        "critica": 8,
        "atencao": 5,
        "regular": 2,
    }.get(general_status, 0)
    if general_status_points:
        reasons.append(
            build_priority_reason(
                code=f"status_geral_{general_status}",
                label="Situação geral da inspeção",
                impact="inspecao",
                points=general_status_points,
                description=f"Situação geral visual/operacional registrada como {format_display_value(general_status)}.",
            )
        )

    problem_distribution = inspection_context.get("problem_distribution")
    distribution_points = {
        "generalizado": 7,
        "espalhado": 5,
        "reboleiras": 3,
        "localizado": 1,
    }.get(problem_distribution, 0)
    if distribution_points:
        reasons.append(
            build_priority_reason(
                code=f"distribuicao_problema_{problem_distribution}",
                label="Distribuição do problema observado",
                impact="inspecao",
                points=distribution_points,
                description=f"Distribuição observada no talhão: {format_display_value(problem_distribution)}.",
            )
        )

    if inspection_context.get("return_needed") is True:
        return_days = inspection_context.get("return_days")
        return_description = "A inspeção indicou necessidade de retorno ao talhão."

        if return_days is not None:
            return_description = (
                f"A inspeção indicou necessidade de retorno em {return_days} dia(s)."
            )

        reasons.append(
            build_priority_reason(
                code="retorno_inspecao_necessario",
                label="Retorno de inspeção necessário",
                impact="inspecao",
                points=5,
                description=return_description,
            )
        )

    visual_severity = inspection_context.get("visual_severity")
    visual_severity_points = {
        "alta": 7,
        "media": 4,
        "baixa": 2,
    }.get(visual_severity, 0)
    if visual_severity_points:
        reasons.append(
            build_priority_reason(
                code=f"severidade_visual_{visual_severity}",
                label="Severidade visual",
                impact="inspecao",
                points=visual_severity_points,
                description=f"Severidade visual registrada como {format_display_value(visual_severity)}.",
            )
        )

    defoliation_level = inspection_context.get("defoliation_level")
    defoliation_points = {
        "alta": 5,
        "media": 3,
        "baixa": 1,
    }.get(defoliation_level, 0)
    if defoliation_points:
        reasons.append(
            build_priority_reason(
                code=f"desfolha_{defoliation_level}",
                label="Desfolha observada",
                impact="inspecao",
                points=defoliation_points,
                description=f"Desfolha registrada como {format_display_value(defoliation_level)}.",
            )
        )

    action_taken = inspection_context.get("action_taken")
    action_points = {
        "manejo_realizado": 4,
        "consultar_responsavel": 5,
        "monitorar": 2,
    }.get(action_taken, 0)
    if action_points:
        reasons.append(
            build_priority_reason(
                code=f"acao_tomada_{action_taken}",
                label="Ação tomada",
                impact="inspecao",
                points=action_points,
                description=f"Última inspeção registrou ação tomada: {format_display_value(action_taken)}.",
            )
        )

    points, reasons = cap_priority_reasons(
        reasons,
        PRIORITY_BLOCK_LIMITS["inspection"],
    )

    return {
        "points": points,
        "reasons": reasons,
        "days_since_inspection": days_since_inspection,
    }


def calculate_operational_priority(
    *,
    field: dict,
    spray_context: dict,
    inspection_context: dict,
    calendar_summary: dict,
    current_date: date,
) -> dict:
    reasons = []
    overdue_events_count = len(calendar_summary["overdue_events"])
    overdue_reapplications_count = len(calendar_summary["overdue_reapplications"])
    today_events_count = len(calendar_summary["today_events"])
    active_alerts_count = len(calendar_summary["active_alerts"])

    if overdue_events_count:
        reasons.append(
            build_priority_reason(
                code="evento_calendario_vencido",
                label="Evento vencido no calendário",
                impact="operacional",
                points=min(7, 4 + overdue_events_count),
                description=f"{overdue_events_count} evento(s) previsto(s) estão vencidos.",
            )
        )

    if overdue_reapplications_count:
        reasons.append(
            build_priority_reason(
                code="reaplicacao_prevista_vencida",
                label="Reaplicação prevista vencida",
                impact="operacional",
                points=min(8, 5 + overdue_reapplications_count),
                description=f"{overdue_reapplications_count} reaplicação(ões) prevista(s) estão vencidas.",
            )
        )

    if today_events_count:
        reasons.append(
            build_priority_reason(
                code="evento_previsto_hoje",
                label="Evento previsto para hoje",
                impact="operacional",
                points=5,
                description=f"{today_events_count} evento(s) previsto(s) para hoje.",
            )
        )

    if active_alerts_count:
        reasons.append(
            build_priority_reason(
                code="alerta_operacional_ativo",
                label="Alerta operacional ativo",
                impact="operacional",
                points=min(5, 2 + active_alerts_count),
                description=f"{active_alerts_count} alerta(s) operacional(is) ativo(s) no calendário.",
            )
        )

    if inspection_context.get("pests_found") is True:
        pest_notes = inspection_context.get("pest_notes")
        description = "A inspeção registrou presença visual de pragas."

        if pest_notes:
            description = f"A inspeção registrou presença visual de pragas: {pest_notes}."

        reasons.append(
            build_priority_reason(
                code="pragas_observadas",
                label="Pragas observadas",
                impact="operacional",
                points=3,
                description=description,
            )
        )

    if inspection_context.get("weeds_found") is True:
        weed_pressure = inspection_context.get("weed_pressure")
        description = "A inspeção registrou presença visual de plantas daninhas."

        if weed_pressure:
            description = (
                "A inspeção registrou presença visual de plantas daninhas "
                f"com pressão {format_display_value(weed_pressure).lower()}."
            )

        reasons.append(
            build_priority_reason(
                code="plantas_daninhas_observadas",
                label="Plantas daninhas observadas",
                impact="operacional",
                points=3,
                description=description,
            )
        )

    days_since_spray = spray_context.get("days_since_application")
    days_since_inspection = days_since_optional_datetime(
        inspection_context.get("last_inspection_date")
        or inspection_context.get("inspected_at"),
        current_date,
    )
    has_recent_spray = (
        isinstance(days_since_spray, int)
        and days_since_spray <= PRIORITY_RECENT_RECORD_WINDOW_DAYS
    )
    has_recent_inspection = (
        isinstance(days_since_inspection, int)
        and days_since_inspection <= PRIORITY_RECENT_RECORD_WINDOW_DAYS
    )

    if (
        field.get("status_lavoura") == "em_campo"
        and not has_recent_spray
        and not has_recent_inspection
    ):
        reasons.append(
            build_priority_reason(
                code="talhao_ativo_sem_registro_recente",
                label="Talhão ativo sem registro recente",
                impact="operacional",
                points=8,
                description="Talhão ativo sem pulverização ou inspeção recente registrada.",
            )
        )

    points, reasons = cap_priority_reasons(
        reasons,
        PRIORITY_BLOCK_LIMITS["operational"],
    )

    return {
        "points": points,
        "reasons": reasons,
        "overdue_events_count": overdue_events_count,
        "overdue_reapplications_count": overdue_reapplications_count,
        "today_events_count": today_events_count,
        "active_alerts_count": active_alerts_count,
        "active_without_recent_record": any(
            reason["code"] == "talhao_ativo_sem_registro_recente"
            for reason in reasons
        ),
    }


def calculate_history_priority(field: dict) -> dict:
    reasons = []

    historical_pressure = field.get("historical_pressure")
    pressure_points = {"alta": 6, "media": 4, "baixa": 1}.get(
        historical_pressure,
        0,
    )
    if pressure_points:
        reasons.append(
            build_priority_reason(
                code=f"pressao_historica_{historical_pressure}",
                label="Pressão histórica",
                impact="historico",
                points=pressure_points,
                description=f"Pressão histórica informada como {format_display_value(historical_pressure)}.",
            )
        )

    peanut_repetition_years = field.get("peanut_repetition_years")
    if peanut_repetition_years is not None:
        if peanut_repetition_years >= 3:
            points = 4
        elif peanut_repetition_years == 2:
            points = 2
        else:
            points = 0

        if points:
            reasons.append(
                build_priority_reason(
                    code="repeticao_amendoim",
                    label="Repetição de amendoim",
                    impact="historico",
                    points=points,
                    description=f"Área com {peanut_repetition_years} safra(s) recente(s) de amendoim.",
                )
            )

    if field.get("crop_rotation") is False:
        reasons.append(
            build_priority_reason(
                code="sem_rotacao_cultura",
                label="Sem rotação de cultura",
                impact="historico",
                points=3,
                description="Não há rotação de cultura informada para o talhão.",
            )
        )

    disease_incidence_level = field.get("disease_incidence_level")
    incidence_points = {"alta": 5, "media": 3, "baixa": 1}.get(
        disease_incidence_level,
        0,
    )
    if incidence_points:
        reasons.append(
            build_priority_reason(
                code=f"historico_problema_{disease_incidence_level}",
                label="Histórico recorrente de problema",
                impact="historico",
                points=incidence_points,
                description=f"Incidência histórica informada como {format_display_value(disease_incidence_level)}.",
            )
        )
    elif field.get("had_disease_incidence") is True:
        reasons.append(
            build_priority_reason(
                code="historico_problema_registrado",
                label="Histórico recorrente de problema",
                impact="historico",
                points=3,
                description="Há registro anterior de incidência no talhão.",
            )
        )

    points, reasons = cap_priority_reasons(
        reasons,
        PRIORITY_BLOCK_LIMITS["history"],
    )

    return {
        "points": points,
        "reasons": reasons,
    }


def calculate_risk_priority(analysis: dict) -> dict:
    reasons = []
    risk = analysis["risk"]
    risk_classification = risk["classification"]
    normalized_risk = normalize_match_text(risk_classification).upper()

    classification_points = {
        "CRITICO": 5,
        "ALTO": 4,
        "MODERADO": 2,
        "BAIXO": 0,
    }.get(normalized_risk, 0)
    if classification_points:
        reasons.append(
            build_priority_reason(
                code=f"risco_classificacao_{normalized_risk.lower()}",
                label="Classificação de risco",
                impact="risco",
                points=classification_points,
                description=f"Classificação de risco atual: {format_display_value(risk_classification)}.",
            )
        )

    agronomic_index = risk.get("agronomic_index") or 0
    if agronomic_index >= 80:
        agronomic_points = 3
    elif agronomic_index >= 60:
        agronomic_points = 2
    elif agronomic_index >= 40:
        agronomic_points = 1
    else:
        agronomic_points = 0

    if agronomic_points:
        reasons.append(
            build_priority_reason(
                code="indice_agronomico_atual",
                label="Risco climático/agronômico",
                impact="risco",
                points=agronomic_points,
                description=f"Índice agronômico atual em {agronomic_index}.",
            )
        )

    climate_index = risk.get("climate_index") or 0
    if climate_index >= 80:
        reasons.append(
            build_priority_reason(
                code="indice_climatico_elevado",
                label="Risco climático/agronômico",
                impact="risco",
                points=1,
                description=f"Índice climático atual em {climate_index}.",
            )
        )

    crop_stage = analysis["field"].get("crop_stage", "")
    normalized_stage = normalize_match_text(crop_stage)
    if any(
        term in normalized_stage
        for term in ["critica", "florescimento", "pegamento", "enchimento"]
    ):
        reasons.append(
            build_priority_reason(
                code="estagio_sensivel_cultura",
                label="Estágio sensível da cultura",
                impact="risco",
                points=2,
                description=f"Estágio informado: {crop_stage}.",
            )
        )

    points, reasons = cap_priority_reasons(
        reasons,
        PRIORITY_BLOCK_LIMITS["risk"],
    )

    return {
        "points": points,
        "reasons": reasons,
    }


def calculate_priority_confidence(
    *,
    field: dict,
    risk_context: dict,
    spray_context: dict,
    inspection_context: dict,
    calendar_events: list[dict],
) -> dict:
    score = 0
    components = {
        "risk_data": False,
        "spray_data": False,
        "inspection_data": False,
        "calendar_data": False,
        "history_data": False,
        "field_data": False,
    }

    if risk_context.get("risk_classification") and risk_context.get(
        "agronomic_index"
    ) is not None:
        score += 25
        components["risk_data"] = True

    if spray_context.get("has_spray_record"):
        score += 20
        components["spray_data"] = True

    if inspection_context.get("has_inspection_record"):
        score += 20
        components["inspection_data"] = True

    if calendar_events:
        score += 10
        components["calendar_data"] = True

    if any(
        field.get(key) is not None
        for key in [
            "historical_pressure",
            "peanut_repetition_years",
            "crop_rotation",
            "disease_incidence_level",
            "had_disease_incidence",
        ]
    ):
        score += 15
        components["history_data"] = True

    if field.get("data_plantio") and field.get("status_lavoura"):
        score += 10
        components["field_data"] = True

    score = min(score, 100)

    return {
        "confidence_score": score,
        "confidence_label": confidence_label_from_score(score),
        "confidence_metrics": {
            **components,
            "calendar_events_count": len(calendar_events),
        },
    }


def build_attention_priority(
    *,
    field: dict,
    analysis: dict,
    spray_context: dict,
    inspection_context: dict,
    calendar_events: list[dict] | None = None,
    current_date: date | None = None,
) -> dict:
    current_date = current_date or datetime.now(ZoneInfo(TIMEZONE)).date()
    calendar_events = calendar_events or get_field_calendar_events(field["id"])
    calendar_summary = summarize_field_calendar_events(
        calendar_events,
        current_date,
    )
    risk = analysis["risk"]
    risk_context = {
        "main_disease": analysis["disease"],
        "risk_classification": risk["classification"],
        "agronomic_index": risk["agronomic_index"],
        "climate_index": risk["climate_index"],
        "crop_stage": analysis["field"].get("crop_stage"),
    }

    defense_priority = calculate_defense_priority(spray_context)
    inspection_priority = calculate_inspection_priority(
        inspection_context,
        current_date,
    )
    operational_priority = calculate_operational_priority(
        field=field,
        spray_context=spray_context,
        inspection_context=inspection_context,
        calendar_summary=calendar_summary,
        current_date=current_date,
    )
    history_priority = calculate_history_priority(field)
    risk_priority = calculate_risk_priority(analysis)
    block_scores = {
        "defense": defense_priority["points"],
        "inspection": inspection_priority["points"],
        "operational": operational_priority["points"],
        "history": history_priority["points"],
        "risk": risk_priority["points"],
    }
    priority_score = min(sum(block_scores.values()), 100)

    if inspection_context.get("symptoms_found") is True:
        priority_score = max(priority_score, 25)

    priority_label = priority_label_from_score(priority_score)
    main_reasons = sorted(
        [
            *defense_priority["reasons"],
            *inspection_priority["reasons"],
            *operational_priority["reasons"],
            *history_priority["reasons"],
            *risk_priority["reasons"],
        ],
        key=lambda reason: reason["points"],
        reverse=True,
    )
    confidence = calculate_priority_confidence(
        field=field,
        risk_context=risk_context,
        spray_context=spray_context,
        inspection_context=inspection_context,
        calendar_events=calendar_events,
    )

    return {
        "priority_score": priority_score,
        "priority_label": priority_label,
        "confidence_score": confidence["confidence_score"],
        "confidence_label": confidence["confidence_label"],
        "main_reasons": main_reasons[:8],
        "metrics": {
            "block_scores": block_scores,
            "defense": {
                "has_spray_record": spray_context.get("has_spray_record"),
                "defense_status": spray_context.get("defense_status"),
                "estimated_defense_percent": spray_context.get(
                    "estimated_defense_percent"
                ),
                "interval_status": spray_context.get("interval_status"),
                "days_since_application": spray_context.get(
                    "days_since_application"
                ),
                "planned_interval_days": spray_context.get(
                    "planned_interval_days"
                ),
                "days_until_reapplication": spray_context.get(
                    "days_until_reapplication"
                ),
            },
            "inspection": {
                "has_inspection_record": inspection_context.get(
                    "has_inspection_record"
                ),
                "days_since_inspection": inspection_priority[
                    "days_since_inspection"
                ],
                "symptoms_found": inspection_context.get("symptoms_found"),
                "visual_severity": inspection_context.get("visual_severity"),
                "defoliation_level": inspection_context.get(
                    "defoliation_level"
                ),
                "action_taken": inspection_context.get("action_taken"),
                "general_status": inspection_context.get("general_status"),
                "problem_distribution": inspection_context.get(
                    "problem_distribution"
                ),
                "pests_found": inspection_context.get("pests_found"),
                "pest_notes": inspection_context.get("pest_notes"),
                "weeds_found": inspection_context.get("weeds_found"),
                "weed_pressure": inspection_context.get("weed_pressure"),
                "soil_condition": inspection_context.get("soil_condition"),
                "return_needed": inspection_context.get("return_needed"),
                "return_days": inspection_context.get("return_days"),
                "observed_area": inspection_context.get("observed_area"),
                "responsible": inspection_context.get("responsible"),
            },
            "operational": {
                "overdue_events_count": operational_priority[
                    "overdue_events_count"
                ],
                "overdue_reapplications_count": operational_priority[
                    "overdue_reapplications_count"
                ],
                "today_events_count": operational_priority["today_events_count"],
                "active_alerts_count": operational_priority[
                    "active_alerts_count"
                ],
                "active_without_recent_record": operational_priority[
                    "active_without_recent_record"
                ],
            },
            "history": {
                "historical_pressure": field.get("historical_pressure"),
                "peanut_repetition_years": field.get("peanut_repetition_years"),
                "crop_rotation": field.get("crop_rotation"),
                "disease_incidence_level": field.get(
                    "disease_incidence_level"
                ),
                "had_disease_incidence": field.get("had_disease_incidence"),
            },
            "risk": risk_context,
            "confidence": confidence["confidence_metrics"],
        },
    }


def disease_matches_target(main_disease: str, target: str | None) -> bool:
    normalized_disease = normalize_match_text(main_disease)
    normalized_target = normalize_match_text(target)

    if not normalized_disease or not normalized_target:
        return False

    return (
        normalized_target in normalized_disease
        or normalized_disease in normalized_target
    )


def get_latest_matching_record(
    records: list[dict],
    date_key: str,
    main_disease: str,
    disease_key: str,
) -> dict | None:
    if not records:
        return None

    sorted_records = sorted(
        records,
        key=lambda record: record[date_key],
        reverse=True,
    )
    matching_records = [
        record
        for record in sorted_records
        if disease_matches_target(main_disease, record.get(disease_key))
    ]

    return (matching_records or sorted_records)[0]


def resolve_defense_reference(
    spray_application: dict,
    selected_product: dict | None = None,
) -> dict:
    planned_interval_days = spray_application.get("planned_interval_days")
    product_default_defense_days = spray_application.get(
        "product_default_defense_days"
    )

    if product_default_defense_days is None and selected_product is not None:
        product_default_defense_days = selected_product.get("default_defense_days")

    if planned_interval_days:
        return {
            "product_default_defense_days": product_default_defense_days,
            "defense_reference_days": planned_interval_days,
            "defense_reference_source": "intervalo_planejado",
        }

    if product_default_defense_days:
        return {
            "product_default_defense_days": product_default_defense_days,
            "defense_reference_days": product_default_defense_days,
            "defense_reference_source": "produto",
        }

    return {
        "product_default_defense_days": None,
        "defense_reference_days": defense_estimator.DEFAULT_PLANNED_INTERVAL_DAYS,
        "defense_reference_source": "fallback",
    }


def resolve_spray_context_reapplication(
    spray_application: dict,
    application_datetime: datetime,
    now: datetime | None = None,
) -> dict:
    reapplication_date = None

    if spray_application.get("reapplication_date"):
        reapplication_date = resolve_spray_reapplication_date(
            spray_application,
            application_datetime,
        )
    elif spray_application.get("generate_reapplication"):
        interval_days = (
            spray_application.get("reapplication_interval_days")
            or spray_application.get("planned_interval_days")
        )

        if interval_days:
            reapplication_date = application_datetime.date() + timedelta(
                days=interval_days
            )

    if reapplication_date is None:
        return {
            "reapplication_date": None,
            "days_until_reapplication": None,
        }

    current_time = normalize_datetime_to_local_timezone(
        now or datetime.now(ZoneInfo(TIMEZONE))
    )

    return {
        "reapplication_date": reapplication_date.isoformat(),
        "days_until_reapplication": (
            reapplication_date - current_time.date()
        ).days,
    }


def build_spray_context(spray_application: dict | None) -> dict:
    if spray_application is None:
        return {
            "has_spray_record": False,
            "last_application_date": None,
            "product_id": None,
            "product": None,
            "product_type": None,
            "product_default_defense_days": None,
            "defense_reference_days": defense_estimator.DEFAULT_PLANNED_INTERVAL_DAYS,
            "defense_reference_source": "fallback",
            "target": None,
            "dose": None,
            "planned_interval_days": None,
            "generate_reapplication": False,
            "reapplication_date": None,
            "days_until_reapplication": None,
            "days_since_application": None,
            "interval_status": "sem_registro",
            "estimated_defense_percent": None,
            "defense_status": "sem_registro",
            "product_target_consistency": "nao_avaliado",
        }

    enriched_spray = enrich_spray_application_interval(spray_application)
    selected_product = None

    if enriched_spray.get("product_id"):
        selected_product = products_repository.get_product(enriched_spray["product_id"])

    defense_reference = resolve_defense_reference(
        enriched_spray,
        selected_product=selected_product,
    )
    application_datetime = datetime.fromisoformat(enriched_spray["application_date"])
    reapplication = resolve_spray_context_reapplication(
        enriched_spray,
        application_datetime,
    )
    defense = defense_estimator.calculate_estimated_defense(
        enriched_spray.get("days_since_application"),
        defense_reference["defense_reference_days"],
    )

    return {
        "has_spray_record": True,
        "last_application_date": enriched_spray["application_date"],
        "product_id": enriched_spray.get("product_id"),
        "product": enriched_spray["product"],
        "product_type": enriched_spray.get("product_type")
        or (selected_product.get("product_type") if selected_product else None),
        "target": enriched_spray["target"],
        "dose": enriched_spray["dose"],
        "planned_interval_days": enriched_spray.get("planned_interval_days"),
        "generate_reapplication": enriched_spray.get("generate_reapplication", False),
        **defense_reference,
        **reapplication,
        "days_since_application": enriched_spray["days_since_application"],
        "interval_status": enriched_spray["interval_status"],
        **defense,
        "product_target_consistency": "nao_avaliado",
    }


def build_inspection_context(inspection: dict | None) -> dict:
    if inspection is None:
        return {
            "has_inspection_record": False,
            "last_inspection_date": None,
            "inspected_at": None,
            "disease": None,
            "symptoms_found": None,
            "visual_severity": None,
            "defoliation_level": None,
            "action_taken": None,
            "general_status": None,
            "problem_distribution": None,
            "pests_found": None,
            "pest_notes": None,
            "weeds_found": None,
            "weed_pressure": None,
            "soil_condition": None,
            "return_needed": None,
            "return_days": None,
            "observed_area": None,
            "responsible": None,
        }

    return {
        "has_inspection_record": True,
        "last_inspection_date": inspection["inspected_at"],
        "inspected_at": inspection["inspected_at"],
        "disease": inspection.get("disease"),
        "symptoms_found": inspection["symptoms_found"],
        "visual_severity": inspection["visual_severity"],
        "defoliation_level": inspection["defoliation_level"],
        "action_taken": inspection.get("action_taken"),
        "general_status": inspection.get("general_status"),
        "problem_distribution": inspection.get("problem_distribution"),
        "pests_found": inspection.get("pests_found"),
        "pest_notes": inspection.get("pest_notes"),
        "weeds_found": inspection.get("weeds_found"),
        "weed_pressure": inspection.get("weed_pressure"),
        "soil_condition": inspection.get("soil_condition"),
        "return_needed": inspection.get("return_needed"),
        "return_days": inspection.get("return_days"),
        "observed_area": inspection.get("observed_area"),
        "responsible": inspection.get("responsible"),
    }


def determine_current_situation(
    field: dict,
    risk_classification: str,
    spray_context: dict,
    inspection_context: dict,
) -> str:
    defense_status = spray_context["defense_status"]

    if field["status_lavoura"] != "em_campo":
        return "sem_prioridade_operacional"

    if is_critical_or_high_risk(risk_classification) and defense_status in {
        "vencida",
        "sem_registro",
    }:
        return "prioridade_maxima"

    if is_critical_or_high_risk(risk_classification) and defense_status == "baixa":
        return "alta_atencao"

    if is_critical_or_high_risk(risk_classification) and defense_status in {
        "alta",
        "muito_alta",
    }:
        return "monitorar_resposta"

    if is_moderate_risk(risk_classification) and inspection_context[
        "symptoms_found"
    ]:
        return "monitorar"

    if is_low_risk(risk_classification) and spray_context["interval_status"] == "em_dia":
        return "estavel"

    return "monitorar"


def build_situation_reasons(
    field: dict,
    main_analysis: dict,
    spray_context: dict,
    inspection_context: dict,
) -> list[str]:
    risk = main_analysis["risk"]
    reasons = [
        f"Referência de risco atual com classificação {format_display_value(risk['classification'])} e índice agronômico {risk['agronomic_index']}.",
    ]

    if field["status_lavoura"] != "em_campo":
        reasons.append("Talhão fora de campo, sem prioridade operacional imediata.")

    if spray_context["has_spray_record"]:
        reasons.append(
            f"Defesa estimada em {spray_context['estimated_defense_percent']}% com status {format_display_value(spray_context['defense_status']).lower()}."
        )
    else:
        reasons.append("Sem pulverização registrada para apoiar a defesa estimada.")

    if inspection_context["has_inspection_record"]:
        symptoms_label = "com sintomas" if inspection_context["symptoms_found"] else "sem sintomas"
        reasons.append(f"Última inspeção registrada {symptoms_label}.")
    else:
        reasons.append("Sem inspeção registrada para confirmar sintomas em campo.")

    return reasons


def build_current_field_situation(field: dict) -> dict:
    if not field.get("doencas_monitoradas"):
        raise HTTPException(
            status_code=400,
            detail="Talhao sem doencas monitoradas para calcular situacao atual",
        )

    try:
        analyses = run_compact_analyses_for_field(field)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error

    if not analyses:
        raise HTTPException(
            status_code=400,
            detail="Nao foi possivel calcular analise de risco do talhao",
        )

    main_analysis = analyses[0]
    risk = main_analysis["risk"]
    main_disease = main_analysis["disease"]

    spray_applications = spray_applications_repository.list_spray_applications_by_field(
        field["id"]
    )
    selected_spray_application = get_latest_matching_record(
        spray_applications,
        date_key="application_date",
        main_disease=main_disease,
        disease_key="target",
    )
    spray_context = build_spray_context(selected_spray_application)

    inspections = inspections_repository.list_inspections_by_field(field["id"])
    selected_inspection = get_latest_matching_record(
        inspections,
        date_key="inspected_at",
        main_disease=main_disease,
        disease_key="disease",
    )
    inspection_context = build_inspection_context(selected_inspection)

    current_situation = determine_current_situation(
        field=field,
        risk_classification=risk["classification"],
        spray_context=spray_context,
        inspection_context=inspection_context,
    )
    reasons = build_situation_reasons(
        field=field,
        main_analysis=main_analysis,
        spray_context=spray_context,
        inspection_context=inspection_context,
    )
    attention_priority = build_attention_priority(
        field=field,
        analysis=main_analysis,
        spray_context=spray_context,
        inspection_context=inspection_context,
    )

    return {
        "field_id": field["id"],
        "field_name": field["nome"],
        "generated_at": datetime.now(ZoneInfo(TIMEZONE)).isoformat(
            timespec="seconds"
        ),
        "current_situation": current_situation,
        "situation_label": SITUATION_LABELS[current_situation],
        "summary": (
            f"{attention_priority['priority_label']} do talhão, "
            f"com defesa estimada {format_display_value(spray_context['defense_status']).lower()} "
            f"e classificação de risco {format_display_value(risk['classification']).lower()}."
        ),
        "risk_context": {
            "main_disease": main_disease,
            "risk_classification": risk["classification"],
            "agronomic_index": risk["agronomic_index"],
            "climate_index": risk["climate_index"],
        },
        "spray_context": spray_context,
        "inspection_context": inspection_context,
        "recommended_next_action": SITUATION_ACTIONS[current_situation],
        "reasons": reasons,
        **attention_priority,
    }


@app.get("/talhoes/{field_id}/situacao-atual")
def get_current_field_situation(field_id: str) -> dict:
    field = fields_repository.get_field(field_id)

    if field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    return build_current_field_situation(field)


SEASON_SITUATION_ORDER = [
    "estavel",
    "monitorar",
    "monitorar_resposta",
    "alta_atencao",
    "prioridade_maxima",
    "sem_prioridade_operacional",
]
SEASON_DEFENSE_ORDER = [
    "muito_alta",
    "alta",
    "media",
    "baixa",
    "vencida",
    "sem_registro",
]
LOW_OR_EXPIRED_DEFENSE_STATUSES = {"baixa", "vencida", "sem_registro"}
CRITICAL_SITUATIONS = {"alta_atencao", "prioridade_maxima"}
RECENT_INSPECTION_WINDOW_DAYS = 7
UPCOMING_EVENT_WINDOW_DAYS = 14


def parse_optional_datetime(value: str | None) -> datetime | None:
    if not value:
        return None

    try:
        return normalize_datetime_to_local_timezone(datetime.fromisoformat(value))
    except ValueError:
        return None


def get_latest_record(records: list[dict], date_key: str) -> dict | None:
    dated_records = [
        record for record in records if record.get(date_key)
    ]

    if not dated_records:
        return None

    return sorted(
        dated_records,
        key=lambda record: record[date_key],
        reverse=True,
    )[0]


def parse_report_timeline_datetime(value: str | date | datetime | None) -> datetime | None:
    if value is None:
        return None

    if isinstance(value, datetime):
        return normalize_datetime_to_local_timezone(value)

    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time()).replace(
            tzinfo=ZoneInfo(TIMEZONE)
        )

    try:
        parsed_value = datetime.fromisoformat(value)
    except ValueError:
        try:
            parsed_date = date.fromisoformat(value)
        except ValueError:
            return None

        return datetime.combine(parsed_date, datetime.min.time()).replace(
            tzinfo=ZoneInfo(TIMEZONE)
        )

    return normalize_datetime_to_local_timezone(parsed_value)


def select_technical_report_field_data(field: dict) -> dict:
    field_keys = [
        "id",
        "nome",
        "cidade",
        "cultura",
        "data_plantio",
        "status_lavoura",
        "doencas_monitoradas",
        "latitude",
        "longitude",
        "previous_crop",
        "crop_rotation",
        "peanut_repetition_years",
        "had_disease_incidence",
        "previous_diseases",
        "disease_incidence_level",
        "historical_pressure",
        "agronomic_history_notes",
    ]

    return {key: field.get(key) for key in field_keys}


def select_latest_inspection_report_data(inspection: dict | None) -> dict | None:
    if inspection is None:
        return None

    inspection_keys = [
        "id",
        "field_id",
        "field_name",
        "disease",
        "symptoms_found",
        "visual_severity",
        "defoliation_level",
        "action_taken",
        "notes",
        "inspected_at",
        "general_status",
        "problem_distribution",
        "pests_found",
        "pest_notes",
        "weeds_found",
        "weed_pressure",
        "soil_condition",
        "return_needed",
        "return_days",
        "observed_area",
        "responsible",
    ]

    return {key: inspection.get(key) for key in inspection_keys}


def select_latest_spray_report_data(
    spray_application: dict | None,
) -> dict | None:
    if spray_application is None:
        return None

    enriched_spray_application = enrich_spray_application_interval(spray_application)
    spray_keys = [
        "id",
        "field_id",
        "field_name",
        "product",
        "target",
        "dose",
        "responsible",
        "planned_interval_days",
        "notes",
        "application_date",
        "days_since_application",
        "interval_status",
    ]

    return {
        key: enriched_spray_application.get(key)
        for key in spray_keys
    }


def build_inspections_summary(inspections: list[dict]) -> dict:
    latest_inspection = get_latest_record(inspections, "inspected_at")

    return {
        "total": len(inspections),
        "latest_date": latest_inspection.get("inspected_at")
        if latest_inspection
        else None,
        "symptoms_found_count": len(
            [
                inspection
                for inspection in inspections
                if inspection.get("symptoms_found") is True
            ]
        ),
        "return_needed_count": len(
            [
                inspection
                for inspection in inspections
                if inspection.get("return_needed") is True
            ]
        ),
        "last_responsible": latest_inspection.get("responsible")
        if latest_inspection
        else None,
        "last_general_status": latest_inspection.get("general_status")
        if latest_inspection
        else None,
        "last_problem_distribution": latest_inspection.get(
            "problem_distribution"
        )
        if latest_inspection
        else None,
    }


def build_spray_summary(spray_applications: list[dict]) -> dict:
    latest_spray_application = get_latest_record(
        spray_applications,
        "application_date",
    )
    latest_spray_report_data = select_latest_spray_report_data(
        latest_spray_application
    )

    return {
        "total": len(spray_applications),
        "latest_date": latest_spray_report_data.get("application_date")
        if latest_spray_report_data
        else None,
        "last_product": latest_spray_report_data.get("product")
        if latest_spray_report_data
        else None,
        "last_target": latest_spray_report_data.get("target")
        if latest_spray_report_data
        else None,
        "last_interval_status": latest_spray_report_data.get("interval_status")
        if latest_spray_report_data
        else None,
        "days_since_last_application": latest_spray_report_data.get(
            "days_since_application"
        )
        if latest_spray_report_data
        else None,
    }


def build_technical_report_timeline(
    field: dict,
    inspections: list[dict],
    spray_applications: list[dict],
) -> list[dict]:
    timeline = []

    if field.get("data_plantio"):
        timeline.append(
            {
                "type": "planting",
                "date": field["data_plantio"],
                "title": "Plantio registrado",
                "description": f"Plantio do talhão {field['nome']} registrado no sistema.",
                "metadata": {
                    "field_id": field["id"],
                    "crop": field.get("cultura"),
                    "crop_status": field.get("status_lavoura"),
                },
            }
        )

    for inspection in inspections:
        if not inspection.get("inspected_at"):
            continue

        timeline.append(
            {
                "type": "inspection",
                "date": inspection["inspected_at"],
                "title": "Inspeção registrada",
                "description": (
                    "Registro observacional de campo para apoiar a priorização "
                    f"operacional do talhão {field['nome']}."
                ),
                "metadata": {
                    "inspection_id": inspection.get("id"),
                    "disease": inspection.get("disease"),
                    "symptoms_found": inspection.get("symptoms_found"),
                    "general_status": inspection.get("general_status"),
                    "problem_distribution": inspection.get(
                        "problem_distribution"
                    ),
                    "return_needed": inspection.get("return_needed"),
                    "responsible": inspection.get("responsible"),
                },
            }
        )

    for spray_application in spray_applications:
        if not spray_application.get("application_date"):
            continue

        enriched_spray_application = enrich_spray_application_interval(
            spray_application
        )
        timeline.append(
            {
                "type": "spray",
                "date": enriched_spray_application["application_date"],
                "title": "Pulverização registrada",
                "description": (
                    "Registro operacional de pulverização informado para o talhão "
                    f"{field['nome']}."
                ),
                "metadata": {
                    "spray_application_id": enriched_spray_application.get("id"),
                    "product": enriched_spray_application.get("product"),
                    "target": enriched_spray_application.get("target"),
                    "planned_interval_days": enriched_spray_application.get(
                        "planned_interval_days"
                    ),
                    "interval_status": enriched_spray_application.get(
                        "interval_status"
                    ),
                    "responsible": enriched_spray_application.get(
                        "responsible"
                    ),
                },
            }
        )

    return sorted(
        timeline,
        key=lambda item: parse_report_timeline_datetime(item.get("date"))
        or datetime.min.replace(tzinfo=ZoneInfo(TIMEZONE)),
        reverse=True,
    )


def build_field_technical_report(field: dict) -> dict:
    inspections = inspections_repository.list_inspections_by_field(field["id"])
    spray_applications = (
        spray_applications_repository.list_spray_applications_by_field(field["id"])
    )
    latest_inspection = get_latest_record(inspections, "inspected_at")
    latest_spray_application = get_latest_record(
        spray_applications,
        "application_date",
    )

    return {
        "field": select_technical_report_field_data(field),
        "current_situation": build_current_field_situation(field),
        "latest_inspection": select_latest_inspection_report_data(
            latest_inspection
        ),
        "latest_spray_application": select_latest_spray_report_data(
            latest_spray_application
        ),
        "inspections_summary": build_inspections_summary(inspections),
        "spray_summary": build_spray_summary(spray_applications),
        "timeline": build_technical_report_timeline(
            field,
            inspections,
            spray_applications,
        ),
        "generated_at": datetime.now(ZoneInfo(TIMEZONE)).isoformat(
            timespec="seconds"
        ),
        "safety_note": TECHNICAL_REPORT_SAFETY_NOTE,
    }


@app.get("/talhoes/{field_id}/relatorio-tecnico")
def get_field_technical_report(field_id: str) -> dict:
    field = fields_repository.get_field(field_id)

    if field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    return build_field_technical_report(field)


def select_operational_context_field_data(field: dict) -> dict:
    return {
        "id": field.get("id"),
        "nome": field.get("nome"),
        "name": field.get("nome"),
        "cidade": field.get("cidade"),
        "city": field.get("cidade"),
        "cultura": field.get("cultura"),
        "crop": field.get("cultura"),
        "data_plantio": field.get("data_plantio"),
        "planting_date": field.get("data_plantio"),
        "status_lavoura": field.get("status_lavoura"),
        "crop_status": field.get("status_lavoura"),
        "doencas_monitoradas": field.get("doencas_monitoradas"),
        "monitored_diseases": field.get("doencas_monitoradas"),
        "manual_crop_stage": field.get("manual_crop_stage"),
        "crop_stage_updated_at": field.get("crop_stage_updated_at"),
        "crop_stage_notes": field.get("crop_stage_notes"),
    }


def crop_stage_label(stage: str) -> str:
    labels = {
        "plantio": "Plantio",
        "emergencia_estabelecimento": "Emergência e estabelecimento",
        "vegetativo": "Vegetativo",
        "florescimento": "Florescimento",
        "enchimento_vagens": "Enchimento de vagens",
        "pre_arranquio": "Pré-arranquio",
        "arranquio": "Arranquio",
        "colheita": "Colheita",
        "pos_colheita": "Pós-colheita",
        "nao_informado": "Não informado",
    }

    return labels.get(stage, stage.replace("_", " ").title())


def estimate_crop_stage_from_days(days_after_planting: int | None) -> str:
    if days_after_planting is None or days_after_planting < 0:
        return "nao_informado"

    if days_after_planting <= 7:
        return "plantio"

    if days_after_planting <= 20:
        return "emergencia_estabelecimento"

    if days_after_planting <= 40:
        return "vegetativo"

    if days_after_planting <= 70:
        return "florescimento"

    if days_after_planting <= 100:
        return "enchimento_vagens"

    if days_after_planting <= 120:
        return "pre_arranquio"

    return "arranquio"


def calculate_days_after_planting(planting_date: str | None) -> int | None:
    if not planting_date:
        return None

    try:
        parsed_planting_date = date.fromisoformat(planting_date)
    except ValueError:
        return None

    return (datetime.now(ZoneInfo(TIMEZONE)).date() - parsed_planting_date).days


def build_crop_stage_context(field: dict) -> dict:
    crop_status = field.get("status_lavoura")
    manual_crop_stage = field.get("manual_crop_stage")
    status_stage_map = {
        "pre_arranquio": "pre_arranquio",
        "arrancado": "arranquio",
        "colhido": "pos_colheita",
    }
    days_after_planting = calculate_days_after_planting(field.get("data_plantio"))

    if manual_crop_stage:
        stage = manual_crop_stage
        source = "manual"
    elif crop_status in status_stage_map:
        stage = status_stage_map[crop_status]
        source = "manual"
    elif days_after_planting is None:
        stage = "nao_informado"
        source = "missing"
    else:
        stage = estimate_crop_stage_from_days(days_after_planting)
        source = "estimated"

    descriptions = {
        "plantio": "Leitura inicial do ciclo, baseada no registro de plantio.",
        "emergencia_estabelecimento": "Talhão em estabelecimento inicial, leitura operacional estimada.",
        "vegetativo": "Talhão em desenvolvimento vegetativo, leitura operacional estimada.",
        "florescimento": "Fase sensível do ciclo, útil para priorizar acompanhamento de campo.",
        "enchimento_vagens": "Fase sensível para acompanhar condição do talhão e registros recentes.",
        "pre_arranquio": "Talhão próximo ao arranquio ou em pré-arranquio informado.",
        "arranquio": "Talhão em arranquio ou já arrancado, com leitura operacional reduzida.",
        "pos_colheita": "Talhão em pós-colheita, útil para memória da propriedade.",
        "nao_informado": "Não há dados suficientes para estimar a fase da lavoura.",
    }

    description = descriptions.get(
        stage,
        "Fase informada manualmente para apoiar a leitura operacional do talhao.",
    )

    if source == "manual" and field.get("crop_stage_notes"):
        description = f"{description} Observacao: {field['crop_stage_notes']}"

    return {
        "stage": stage,
        "label": crop_stage_label(stage),
        "source": source,
        "days_after_planting": days_after_planting,
        "description": description,
        "updated_at": field.get("crop_stage_updated_at"),
        "notes": field.get("crop_stage_notes"),
    }


def build_missing_weather_context(source: str, summary: str) -> dict:
    return {
        "available": False,
        "source": source,
        "current_temperature_c": None,
        "current_precipitation_mm": None,
        "current_relative_humidity": None,
        "recent_rain_7d_mm": None,
        "forecast_rain_7d_mm": None,
        "forecast_max_temp_avg_c": None,
        "forecast_min_temp_avg_c": None,
        "reference_et0_7d_mm": None,
        "soil_moisture_latest": None,
        "soil_temperature_latest_c": None,
        "summary": summary,
    }


def parse_coordinate(value) -> float | None:
    if value in (None, ""):
        return None

    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def build_field_weather_context(field: dict) -> dict:
    latitude = parse_coordinate(field.get("latitude"))
    longitude = parse_coordinate(field.get("longitude"))

    if latitude is None or longitude is None:
        return build_missing_weather_context(
            "missing_location",
            "Nao ha localizacao suficiente para cruzar dados meteorologicos deste talhao.",
        )

    try:
        return fetch_weather_context(latitude, longitude)
    except Exception:
        return build_missing_weather_context(
            "unavailable",
            "Dados meteorologicos externos indisponiveis no momento.",
        )


def is_low_rainfall(value: float | int | None) -> bool:
    return value is not None and value < 10


def is_very_low_rainfall(value: float | int | None) -> bool:
    return value is not None and value < 5


def is_significant_rainfall(value: float | int | None) -> bool:
    return value is not None and value >= 20


def refine_water_context_with_weather(
    *,
    status: str,
    main_reason: str,
    farmer_message: str,
    latest_soil_condition: str | None,
    is_sensitive_stage: bool,
    weather: dict,
) -> tuple[str, str, str]:
    if not weather.get("available"):
        if (
            not latest_soil_condition
            and weather.get("source") == "missing_location"
        ):
            farmer_message = (
                "Nao ha localizacao suficiente para cruzar dados meteorologicos deste talhao."
            )

        return status, main_reason, farmer_message

    recent_rain = weather.get("recent_rain_7d_mm")
    forecast_rain = weather.get("forecast_rain_7d_mm")
    forecast_max_temp = weather.get("forecast_max_temp_avg_c")
    reference_et0 = weather.get("reference_et0_7d_mm")
    low_recent_and_forecast = is_low_rainfall(recent_rain) and is_low_rainfall(
        forecast_rain,
    )
    hot_or_high_et0 = (
        (
            forecast_max_temp is not None
            and forecast_max_temp >= 33
        )
        or (
            reference_et0 is not None
            and reference_et0 >= 30
        )
    )

    if latest_soil_condition == "seco":
        if is_very_low_rainfall(forecast_rain) and is_sensitive_stage:
            return (
                "critico",
                "Ultima inspecao registrou solo seco em fase sensivel e a previsao indica pouca chuva.",
                "A ultima inspecao registrou solo seco e a previsao indica pouca chuva nos proximos dias. Vale priorizar acompanhamento de campo.",
            )

        if is_low_rainfall(forecast_rain) or hot_or_high_et0:
            return (
                "atencao",
                "Ultima inspecao registrou solo seco e a meteorologia indica atencao hidrica.",
                "A ultima inspecao registrou solo seco e a previsao indica pouca chuva nos proximos dias. Vale priorizar acompanhamento de campo.",
            )

        if is_significant_rainfall(forecast_rain):
            return (
                "atencao",
                "Ha previsao de chuva, mas a ultima inspecao registrou solo seco.",
                "Ha previsao de chuva nos proximos dias, mas a ultima inspecao registrou condicao de solo desfavoravel. Mantenha o acompanhamento.",
            )

    if latest_soil_condition == "compactado":
        if is_low_rainfall(forecast_rain):
            return (
                "atencao",
                "Ultima inspecao registrou solo compactado e a previsao indica pouca chuva.",
                "A ultima inspecao registrou solo compactado e a previsao indica pouca chuva nos proximos dias. Vale acompanhar a area em campo.",
            )

        if is_significant_rainfall(forecast_rain):
            return (
                "atencao",
                "Ha previsao de chuva, mas a ultima inspecao registrou solo compactado.",
                "Ha previsao de chuva nos proximos dias, mas a ultima inspecao registrou condicao de solo desfavoravel. Mantenha o acompanhamento.",
            )

    if latest_soil_condition in {"adequado", "umido"}:
        if low_recent_and_forecast:
            return (
                "atencao",
                "Chuva recente e previsao dos proximos dias estao baixas.",
                "A condicao visual do solo foi favoravel, mas chuva recente e prevista estao baixas. Mantenha o acompanhamento.",
            )

        if is_significant_rainfall(recent_rain) or is_significant_rainfall(
            forecast_rain,
        ):
            return (
                "adequado",
                "Dados de solo e meteorologia nao indicam restricao hidrica relevante.",
                "Os dados meteorologicos nao indicam restricao hidrica relevante nos proximos dias.",
            )

    if latest_soil_condition is None:
        if low_recent_and_forecast:
            return (
                "atencao",
                "Chuva recente e previsao dos proximos dias estao baixas.",
                "Chuva recente e previsao dos proximos dias estao baixas. Vale verificar a condicao do solo em campo.",
            )

        if is_significant_rainfall(recent_rain) or is_significant_rainfall(
            forecast_rain,
        ):
            return (
                "adequado",
                "Meteorologia indica chuva recente ou prevista suficiente para leitura operacional inicial.",
                "Os dados meteorologicos nao indicam restricao hidrica relevante nos proximos dias.",
            )

    return status, main_reason, farmer_message


def build_water_context(
    latest_inspection: dict | None,
    crop_stage_context: dict,
    field: dict,
) -> dict:
    latest_soil_condition = (
        latest_inspection.get("soil_condition") if latest_inspection else None
    )
    sensitive_stages = {"florescimento", "enchimento_vagens", "pre_arranquio"}
    is_sensitive_stage = crop_stage_context["stage"] in sensitive_stages
    notes = " ".join(
        str(value).lower()
        for value in [
            field.get("agronomic_history_notes"),
            latest_inspection.get("notes") if latest_inspection else None,
        ]
        if value
    )
    weather = build_field_weather_context(field)

    if not latest_soil_condition:
        status = "sem_dados"
        main_reason = "Sem registro recente de condição visual do solo."
        farmer_message = "Não há dados recentes de condição do solo para este talhão."
    elif latest_soil_condition in {"seco", "compactado"}:
        status = "critico" if is_sensitive_stage else "atencao"
        main_reason = f"Última inspeção registrou solo {format_display_value(latest_soil_condition).lower()}."
        farmer_message = (
            "O talhão está em fase sensível e possui registro de condição hídrica desfavorável."
            if is_sensitive_stage
            else "A última inspeção registrou solo seco. Vale priorizar acompanhamento de campo."
        )
    elif latest_soil_condition == "encharcado":
        status = "critico" if is_sensitive_stage else "atencao"
        main_reason = "Última inspeção registrou solo encharcado."
        farmer_message = (
            "A última inspeção registrou condição hídrica desfavorável. Acompanhar a área em campo."
        )
    elif latest_soil_condition in {"adequado", "umido"}:
        status = "adequado"
        main_reason = "Última inspeção registrou condição visual do solo sem alerta hídrico."
        farmer_message = "A última inspeção registrou condição visual do solo adequada ou úmida."
    else:
        status = "sem_dados"
        main_reason = "Condição visual do solo não avaliada."
        farmer_message = "Não há leitura conclusiva de condição do solo para este talhão."

    if status == "adequado" and ("seca" in notes or "seco" in notes):
        status = "atencao"
        main_reason = "Há observações históricas relacionadas à seca ou solo seco."
        farmer_message = "Há registro textual de condição seca. Vale acompanhar a área em campo."

    status, main_reason, farmer_message = refine_water_context_with_weather(
        status=status,
        main_reason=main_reason,
        farmer_message=farmer_message,
        latest_soil_condition=latest_soil_condition,
        is_sensitive_stage=is_sensitive_stage,
        weather=weather,
    )

    return {
        "status": status,
        "label": format_display_value(status),
        "latest_soil_condition": latest_soil_condition,
        "is_sensitive_stage": is_sensitive_stage,
        "main_reason": main_reason,
        "farmer_message": farmer_message,
        "weather": weather,
    }


def normalize_recurrent_diseases(value: str | list[str] | None) -> list[str]:
    if not value:
        return []

    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]

    separators = [",", ";"]
    normalized_values = [value]
    for separator in separators:
        normalized_values = [
            part
            for item in normalized_values
            for part in item.split(separator)
        ]

    return [item.strip() for item in normalized_values if item.strip()]


def build_historical_memory(field: dict) -> dict:
    recurrent_diseases = normalize_recurrent_diseases(field.get("previous_diseases"))
    pressure_level = (
        field.get("historical_pressure")
        or field.get("disease_incidence_level")
        or "sem_dados"
    )
    rotation_attention = (
        field.get("crop_rotation") is False
        or (field.get("peanut_repetition_years") or 0) > 0
    )
    history_values = [
        field.get("previous_crop"),
        field.get("crop_rotation"),
        field.get("peanut_repetition_years"),
        field.get("had_disease_incidence"),
        field.get("previous_diseases"),
        field.get("disease_incidence_level"),
        field.get("historical_pressure"),
        field.get("agronomic_history_notes"),
    ]
    has_history = any(value not in (None, "", []) for value in history_values)
    attention_points = []

    if field.get("had_disease_incidence") is True:
        attention_points.append("Há registro de incidência anterior no talhão.")

    if recurrent_diseases:
        attention_points.append(
            "Doenças recorrentes registradas: " + ", ".join(recurrent_diseases)
        )

    if rotation_attention:
        attention_points.append("Histórico indica atenção à rotação de cultura.")

    if pressure_level in {"media", "alta"}:
        attention_points.append(
            f"Pressão histórica {format_display_value(pressure_level).lower()}."
        )

    summary = (
        "Histórico operacional disponível para apoiar a leitura do talhão."
        if has_history
        else "Sem histórico agronômico cadastrado para este talhão."
    )

    return {
        "has_history": has_history,
        "pressure_level": pressure_level,
        "recurrent_diseases": recurrent_diseases,
        "rotation_attention": rotation_attention,
        "summary": summary,
        "attention_points": attention_points,
    }


def build_field_observation_context(latest_inspection: dict | None) -> dict:
    if latest_inspection is None:
        return {
            "has_recent_inspection": False,
            "latest_inspection_date": None,
            "symptoms_found": None,
            "pests_found": None,
            "weeds_found": None,
            "return_needed": None,
            "summary": "Sem inspeção registrada para consolidar observações de campo.",
            "attention_points": ["Registrar nova inspeção de campo."],
        }

    inspection_datetime = parse_optional_datetime(latest_inspection.get("inspected_at"))
    current_date = datetime.now(ZoneInfo(TIMEZONE)).date()
    has_recent_inspection = (
        inspection_datetime is not None
        and (current_date - inspection_datetime.date()).days
        <= RECENT_INSPECTION_WINDOW_DAYS
    )
    attention_points = []

    if latest_inspection.get("symptoms_found") is True:
        attention_points.append("Sintomas observados na última inspeção.")

    if latest_inspection.get("pests_found") is True:
        attention_points.append("Pragas observadas na última inspeção.")

    if latest_inspection.get("weeds_found") is True:
        attention_points.append("Plantas daninhas observadas na última inspeção.")

    if latest_inspection.get("return_needed") is True:
        attention_points.append("Retorno/reinspeção indicado no registro.")

    summary = (
        "Última inspeção recente disponível para leitura operacional."
        if has_recent_inspection
        else "Há inspeção registrada, mas ela pode não representar a condição atual do talhão."
    )

    return {
        "has_recent_inspection": has_recent_inspection,
        "latest_inspection_date": latest_inspection.get("inspected_at"),
        "symptoms_found": latest_inspection.get("symptoms_found"),
        "pests_found": latest_inspection.get("pests_found"),
        "weeds_found": latest_inspection.get("weeds_found"),
        "return_needed": latest_inspection.get("return_needed"),
        "summary": summary,
        "attention_points": attention_points,
    }


def build_application_response_context(
    latest_spray_application: dict | None,
    inspections: list[dict],
) -> dict:
    if latest_spray_application is None:
        return {
            "has_spray": False,
            "latest_spray_date": None,
            "latest_product": None,
            "has_inspection_after_spray": False,
            "status": "sem_aplicacao",
            "farmer_message": "Não há aplicação registrada para avaliar resposta pós-aplicação.",
        }

    latest_spray_date = parse_optional_datetime(
        latest_spray_application.get("application_date")
    )
    has_inspection_after_spray = False

    if latest_spray_date is not None:
        has_inspection_after_spray = any(
            inspection_datetime is not None and inspection_datetime > latest_spray_date
            for inspection_datetime in [
                parse_optional_datetime(inspection.get("inspected_at"))
                for inspection in inspections
            ]
        )

    if has_inspection_after_spray:
        status = "verificada_por_inspecao"
        farmer_message = "Há inspeção posterior à aplicação para apoiar o acompanhamento da resposta da área."
    else:
        status = "resposta_nao_verificada"
        farmer_message = "Resposta observada ainda não verificada por inspeção posterior. Acompanhar resposta da área."

    return {
        "has_spray": True,
        "latest_spray_date": latest_spray_application.get("application_date"),
        "latest_product": latest_spray_application.get("product"),
        "has_inspection_after_spray": has_inspection_after_spray,
        "status": status,
        "farmer_message": farmer_message,
    }


def build_data_quality_context(
    field: dict,
    inspections: list[dict],
    latest_inspection: dict | None,
    spray_applications: list[dict],
    historical_memory: dict,
    field_observation_context: dict,
) -> dict:
    score = 0
    missing_items = []

    if field.get("data_plantio"):
        score += 15
    else:
        missing_items.append("Data de plantio")

    if latest_inspection:
        score += 20
    else:
        missing_items.append("Inspeção de campo")

    if field_observation_context["has_recent_inspection"]:
        score += 15
    else:
        missing_items.append("Inspeção recente")

    if historical_memory["has_history"]:
        score += 15
    else:
        missing_items.append("Histórico agronômico")

    if latest_inspection and latest_inspection.get("soil_condition"):
        score += 15
    else:
        missing_items.append("Condição visual do solo")

    if spray_applications:
        score += 10
    else:
        missing_items.append("Registro de pulverização")

    if any(inspection.get("responsible") for inspection in inspections):
        score += 10
    else:
        missing_items.append("Responsável nas inspeções")

    if score >= 75:
        label = "alta"
    elif score >= 45:
        label = "media"
    else:
        label = "baixa"

    return {
        "score": score,
        "label": label,
        "missing_items": missing_items,
        "summary": (
            "Dados suficientes para uma boa leitura operacional."
            if label == "alta"
            else "A leitura operacional pode melhorar com mais registros de campo."
        ),
    }


def build_farmer_summary(
    crop_stage_context: dict,
    water_context: dict,
    historical_memory: dict,
    field_observation_context: dict,
    application_response_context: dict,
    data_quality_context: dict,
) -> dict:
    main_points = [
        crop_stage_context["description"],
        water_context["farmer_message"],
    ]
    suggested_follow_up = []

    if field_observation_context["attention_points"]:
        main_points.extend(field_observation_context["attention_points"][:2])

    if historical_memory["attention_points"]:
        main_points.append(historical_memory["attention_points"][0])

    if not field_observation_context["has_recent_inspection"]:
        suggested_follow_up.append("Registrar nova inspeção de campo.")

    if water_context["status"] in {"sem_dados", "atencao", "critico"}:
        suggested_follow_up.append("Verificar condição do solo nos próximos dias.")

    if application_response_context["status"] == "resposta_nao_verificada":
        suggested_follow_up.append("Acompanhar resposta da última aplicação.")

    if data_quality_context["label"] != "alta":
        suggested_follow_up.append(
            "Completar histórico do talhão para aumentar a confiança da leitura."
        )

    if not suggested_follow_up:
        suggested_follow_up.append("Manter registros operacionais atualizados.")

    headline = (
        "Talhão com pontos operacionais para acompanhamento."
        if main_points
        else "Talhão sem alertas operacionais relevantes nos dados disponíveis."
    )

    return {
        "headline": headline,
        "main_points": main_points[:5],
        "suggested_follow_up": suggested_follow_up,
    }


def build_field_operational_context(field: dict) -> dict:
    inspections = inspections_repository.list_inspections_by_field(field["id"])
    spray_applications = (
        spray_applications_repository.list_spray_applications_by_field(field["id"])
    )
    latest_inspection = get_latest_record(inspections, "inspected_at")
    latest_spray_application = get_latest_record(
        spray_applications,
        "application_date",
    )
    crop_stage_context = build_crop_stage_context(field)
    water_context = build_water_context(
        latest_inspection,
        crop_stage_context,
        field,
    )
    historical_memory = build_historical_memory(field)
    field_observation_context = build_field_observation_context(latest_inspection)
    application_response_context = build_application_response_context(
        latest_spray_application,
        inspections,
    )
    data_quality_context = build_data_quality_context(
        field,
        inspections,
        latest_inspection,
        spray_applications,
        historical_memory,
        field_observation_context,
    )
    farmer_summary = build_farmer_summary(
        crop_stage_context,
        water_context,
        historical_memory,
        field_observation_context,
        application_response_context,
        data_quality_context,
    )

    return {
        "field": select_operational_context_field_data(field),
        "crop_stage_context": crop_stage_context,
        "water_context": water_context,
        "historical_memory": historical_memory,
        "field_observation_context": field_observation_context,
        "application_response_context": application_response_context,
        "data_quality_context": data_quality_context,
        "farmer_summary": farmer_summary,
        "generated_at": datetime.now(ZoneInfo(TIMEZONE)).isoformat(
            timespec="seconds"
        ),
        "safety_note": OPERATIONAL_CONTEXT_SAFETY_NOTE,
    }


@app.get("/talhoes/{field_id}/contexto-operacional")
def get_field_operational_context(field_id: str) -> dict:
    field = fields_repository.get_field(field_id)

    if field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    return build_field_operational_context(field)


def normalize_pdf_text(value) -> str:
    if value is None or value == "":
        return "N/A"

    if isinstance(value, bool):
        return "Sim" if value else "Não"

    if isinstance(value, (int, float)):
        return str(value)

    if isinstance(value, list):
        return ", ".join(normalize_pdf_text(item) for item in value) or "N/A"

    return str(value)


def format_pdf_status(value) -> str:
    if value is None or value == "":
        return "N/A"

    labels = {
        "sem_registro": "Sem registro",
        "muito_alta": "Muito alta",
        "alta": "Alta",
        "media": "Média",
        "baixa": "Baixa",
        "nenhuma": "Nenhuma",
        "vencida": "Vencida",
        "atrasado": "Atrasado",
        "atencao": "Atenção",
        "em_dia": "Em dia",
        "critico": "Crítico",
        "moderado": "Moderado",
        "baixo": "Baixo",
        "estavel": "Estável",
        "monitorar": "Monitorar",
        "monitoramento": "Monitoramento",
        "alta_atencao": "Alta atenção",
        "prioridade_maxima": "Prioridade máxima",
        "consultar_responsavel": "Consultar responsável",
        "manejo_realizado": "Manejo realizado",
        "boa": "Boa",
        "regular": "Regular",
        "critica": "Crítica",
        "ausente": "Ausente",
        "localizado": "Localizado",
        "reboleiras": "Reboleiras",
        "espalhado": "Espalhado",
        "generalizado": "Generalizado",
        "seco": "Seco",
        "adequado": "Adequado",
        "umido": "Úmido",
        "encharcado": "Encharcado",
        "compactado": "Compactado",
        "nao_avaliado": "Não avaliado",
        "planting": "Plantio",
        "inspection": "Inspeção",
        "spray": "Pulverização",
        "em_campo": "Em campo",
        "pre_arranquio": "Pré-arranquio",
        "arrancado": "Arrancado",
        "colhido": "Colhido",
    }
    normalized = (
        str(value)
        .strip()
        .lower()
        .replace("-", "_")
        .replace(" ", "_")
    )

    if normalized in labels:
        return labels[normalized]

    return str(value).replace("_", " ").title()


def format_pdf_datetime(value: str | None) -> str:
    if not value:
        return "N/A"

    try:
        if "T" in value:
            parsed_value = datetime.fromisoformat(value)
            return parsed_value.strftime("%d/%m/%Y %H:%M")

        parsed_date = date.fromisoformat(value)
        return parsed_date.strftime("%d/%m/%Y")
    except ValueError:
        return value


def pdf_paragraph(text, style: ParagraphStyle) -> Paragraph:
    return Paragraph(escape(normalize_pdf_text(text)), style)


def build_pdf_key_value_table(
    rows: list[tuple[str, object]],
    styles: dict[str, ParagraphStyle],
) -> Table:
    table_data = [
        [
            pdf_paragraph(label, styles["table_label"]),
            pdf_paragraph(value, styles["table_value"]),
        ]
        for label, value in rows
        if value is not None and value != ""
    ]

    if not table_data:
        table_data = [
            [
                pdf_paragraph("Informação", styles["table_label"]),
                pdf_paragraph("N/A", styles["table_value"]),
            ]
        ]

    table = Table(table_data, colWidths=[5.2 * cm, 11.4 * cm], hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D7DED6")),
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F3F6F2")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )

    return table


def add_pdf_section(
    story: list,
    title: str,
    styles: dict[str, ParagraphStyle],
) -> None:
    story.append(Spacer(1, 0.28 * cm))
    story.append(pdf_paragraph(title, styles["section_title"]))
    story.append(Spacer(1, 0.12 * cm))


def build_pdf_filename(field_name: str) -> str:
    slug = (
        field_name.lower()
        .replace("ã", "a")
        .replace("á", "a")
        .replace("à", "a")
        .replace("â", "a")
        .replace("é", "e")
        .replace("ê", "e")
        .replace("í", "i")
        .replace("ó", "o")
        .replace("ô", "o")
        .replace("õ", "o")
        .replace("ú", "u")
        .replace("ç", "c")
    )
    slug = "".join(character if character.isalnum() else "-" for character in slug)
    slug = "-".join(part for part in slug.split("-") if part)

    return f"relatorio-tecnico-{slug or 'talhao'}.pdf"


def draw_pdf_footer(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#59665D"))
    canvas.drawString(1.7 * cm, 1.1 * cm, "PeanuTec · Relatório operacional de apoio")
    canvas.drawRightString(
        A4[0] - 1.7 * cm,
        1.1 * cm,
        f"Página {document.page}",
    )
    canvas.restoreState()


def build_field_technical_report_pdf(report: dict) -> bytes:
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=1.7 * cm,
        leftMargin=1.7 * cm,
        topMargin=1.6 * cm,
        bottomMargin=1.8 * cm,
        title="Relatório Técnico do Talhão",
    )
    sample_styles = getSampleStyleSheet()
    styles = {
        "title": ParagraphStyle(
            "PeanuTecTitle",
            parent=sample_styles["Title"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#123B2D"),
            spaceAfter=6,
        ),
        "subtitle": ParagraphStyle(
            "PeanuTecSubtitle",
            parent=sample_styles["BodyText"],
            fontName="Helvetica",
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#59665D"),
        ),
        "section_title": ParagraphStyle(
            "PeanuTecSectionTitle",
            parent=sample_styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#123B2D"),
            spaceAfter=4,
        ),
        "body": ParagraphStyle(
            "PeanuTecBody",
            parent=sample_styles["BodyText"],
            fontName="Helvetica",
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#27312B"),
        ),
        "note": ParagraphStyle(
            "PeanuTecNote",
            parent=sample_styles["BodyText"],
            fontName="Helvetica",
            fontSize=9,
            leading=13,
            leftIndent=6,
            rightIndent=6,
            textColor=colors.HexColor("#27312B"),
        ),
        "table_label": ParagraphStyle(
            "PeanuTecTableLabel",
            parent=sample_styles["BodyText"],
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#59665D"),
        ),
        "table_value": ParagraphStyle(
            "PeanuTecTableValue",
            parent=sample_styles["BodyText"],
            fontName="Helvetica",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#27312B"),
        ),
    }
    field = report["field"]
    current_situation = report["current_situation"]
    latest_inspection = report["latest_inspection"]
    latest_spray = report["latest_spray_application"]
    inspections_summary = report["inspections_summary"]
    spray_summary = report["spray_summary"]
    spray_context = current_situation.get("spray_context", {})
    story = [
        pdf_paragraph("PeanuTec", styles["subtitle"]),
        pdf_paragraph("Relatório Técnico do Talhão", styles["title"]),
        pdf_paragraph(
            f"{field.get('nome')} · Gerado em {format_pdf_datetime(report.get('generated_at'))}",
            styles["subtitle"],
        ),
        Spacer(1, 0.22 * cm),
    ]
    note_table = Table(
        [[pdf_paragraph(report["safety_note"], styles["note"])]],
        colWidths=[16.6 * cm],
    )
    note_table.setStyle(
        TableStyle(
            [
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#86A39A")),
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EEF5F2")),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.append(note_table)

    add_pdf_section(story, "Dados do talhão", styles)
    story.append(
        build_pdf_key_value_table(
            [
                ("Nome", field.get("nome")),
                ("Cidade", field.get("cidade")),
                ("Cultura", field.get("cultura")),
                ("Data de plantio", format_pdf_datetime(field.get("data_plantio"))),
                ("Status da lavoura", format_pdf_status(field.get("status_lavoura"))),
                ("Doenças monitoradas", field.get("doencas_monitoradas")),
                ("Cultura anterior", field.get("previous_crop")),
                ("Rotação de cultura", field.get("crop_rotation")),
                ("Repetição de amendoim", field.get("peanut_repetition_years")),
                ("Histórico de incidência", field.get("had_disease_incidence")),
                ("Pressão histórica", format_pdf_status(field.get("historical_pressure"))),
                ("Observações históricas", field.get("agronomic_history_notes")),
            ],
            styles,
        )
    )

    add_pdf_section(story, "Resumo executivo", styles)
    defense_percent = spray_context.get("estimated_defense_percent")
    story.append(
        build_pdf_key_value_table(
            [
                (
                    "Prioridade de Atenção",
                    f"{current_situation.get('priority_score', 'N/A')} · {format_pdf_status(current_situation.get('priority_label') or current_situation.get('situation_label'))}",
                ),
                (
                    "Confiança da análise",
                    f"{current_situation.get('confidence_score', 'N/A')} · {format_pdf_status(current_situation.get('confidence_label'))}",
                ),
                (
                    "Próximo passo operacional",
                    current_situation.get("recommended_next_action"),
                ),
                (
                    "Defesa estimada",
                    "N/A"
                    if defense_percent is None
                    else f"{defense_percent}% · {format_pdf_status(spray_context.get('defense_status'))}",
                ),
            ],
            styles,
        )
    )

    add_pdf_section(story, "Prioridade atual", styles)
    story.append(
        build_pdf_key_value_table(
            [
                ("Score de prioridade", current_situation.get("priority_score")),
                (
                    "Label de prioridade",
                    format_pdf_status(current_situation.get("priority_label")),
                ),
                ("Score de confiança", current_situation.get("confidence_score")),
                (
                    "Label de confiança",
                    format_pdf_status(current_situation.get("confidence_label")),
                ),
            ],
            styles,
        )
    )
    main_reasons = current_situation.get("main_reasons") or []
    if main_reasons:
        story.append(Spacer(1, 0.12 * cm))
        for reason in main_reasons[:8]:
            story.append(
                pdf_paragraph(
                    f"• {reason.get('label')}: {reason.get('description')}",
                    styles["body"],
                )
            )
    else:
        story.append(pdf_paragraph("Sem motivos adicionais informados.", styles["body"]))

    add_pdf_section(story, "Última inspeção", styles)
    if latest_inspection:
        story.append(
            build_pdf_key_value_table(
                [
                    ("Data", format_pdf_datetime(latest_inspection.get("inspected_at"))),
                    ("Responsável", latest_inspection.get("responsible")),
                    ("Doença", latest_inspection.get("disease")),
                    ("Sintomas", latest_inspection.get("symptoms_found")),
                    (
                        "Severidade visual",
                        format_pdf_status(latest_inspection.get("visual_severity")),
                    ),
                    (
                        "Desfolha",
                        format_pdf_status(latest_inspection.get("defoliation_level")),
                    ),
                    (
                        "Situação geral",
                        format_pdf_status(latest_inspection.get("general_status")),
                    ),
                    (
                        "Distribuição do problema",
                        format_pdf_status(latest_inspection.get("problem_distribution")),
                    ),
                    ("Pragas observadas", latest_inspection.get("pests_found")),
                    ("Observações de pragas", latest_inspection.get("pest_notes")),
                    ("Plantas daninhas", latest_inspection.get("weeds_found")),
                    (
                        "Pressão de plantas daninhas",
                        format_pdf_status(latest_inspection.get("weed_pressure")),
                    ),
                    (
                        "Condição do solo",
                        format_pdf_status(latest_inspection.get("soil_condition")),
                    ),
                    ("Necessidade de retorno", latest_inspection.get("return_needed")),
                    (
                        "Prazo de retorno",
                        None
                        if latest_inspection.get("return_days") is None
                        else f"{latest_inspection.get('return_days')} dias",
                    ),
                    ("Área/ponto observado", latest_inspection.get("observed_area")),
                    ("Observações", latest_inspection.get("notes")),
                ],
                styles,
            )
        )
    else:
        story.append(
            pdf_paragraph(
                "Nenhuma inspeção registrada para este talhão.",
                styles["body"],
            )
        )

    add_pdf_section(story, "Última pulverização", styles)
    if latest_spray:
        story.append(
            build_pdf_key_value_table(
                [
                    ("Produto", latest_spray.get("product")),
                    ("Alvo", latest_spray.get("target")),
                    ("Dose", latest_spray.get("dose")),
                    ("Responsável", latest_spray.get("responsible")),
                    ("Data", format_pdf_datetime(latest_spray.get("application_date"))),
                    (
                        "Intervalo planejado",
                        None
                        if latest_spray.get("planned_interval_days") is None
                        else f"{latest_spray.get('planned_interval_days')} dias",
                    ),
                    (
                        "Dias desde aplicação",
                        latest_spray.get("days_since_application"),
                    ),
                    (
                        "Status do intervalo",
                        format_pdf_status(latest_spray.get("interval_status")),
                    ),
                    ("Observações", latest_spray.get("notes")),
                ],
                styles,
            )
        )
    else:
        story.append(
            pdf_paragraph(
                "Nenhuma pulverização registrada para este talhão.",
                styles["body"],
            )
        )

    add_pdf_section(story, "Resumos operacionais", styles)
    story.append(
        build_pdf_key_value_table(
            [
                ("Total de inspeções", inspections_summary.get("total")),
                (
                    "Inspeções com sintomas",
                    inspections_summary.get("symptoms_found_count"),
                ),
                ("Retornos necessários", inspections_summary.get("return_needed_count")),
                ("Total de pulverizações", spray_summary.get("total")),
                ("Último produto", spray_summary.get("last_product")),
                (
                    "Último status de intervalo",
                    format_pdf_status(spray_summary.get("last_interval_status")),
                ),
            ],
            styles,
        )
    )

    story.append(PageBreak())
    add_pdf_section(story, "Linha do tempo operacional", styles)
    timeline = report.get("timeline") or []
    if timeline:
        for item in timeline[:15]:
            story.append(
                pdf_paragraph(
                    f"{format_pdf_datetime(item.get('date'))} · {format_pdf_status(item.get('type'))} · {item.get('title')}",
                    styles["table_label"],
                )
            )
            story.append(pdf_paragraph(item.get("description"), styles["body"]))
            story.append(Spacer(1, 0.12 * cm))
    else:
        story.append(
            pdf_paragraph(
                "Nenhum evento operacional registrado para este talhão.",
                styles["body"],
            )
        )

    document.build(story, onFirstPage=draw_pdf_footer, onLaterPages=draw_pdf_footer)
    buffer.seek(0)

    return buffer.getvalue()


@app.get("/talhoes/{field_id}/relatorio-tecnico/pdf")
def get_field_technical_report_pdf(field_id: str) -> StreamingResponse:
    field = fields_repository.get_field(field_id)

    if field is None:
        raise HTTPException(status_code=404, detail="Talhao nao encontrado")

    report = build_field_technical_report(field)
    pdf_bytes = build_field_technical_report_pdf(report)
    filename = build_pdf_filename(field["nome"])

    return StreamingResponse(
        BytesIO(pdf_bytes),
        media_type="application/pdf",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
    )


def build_fallback_field_situation(field: dict) -> dict:
    spray_application = get_latest_record(
        spray_applications_repository.list_spray_applications_by_field(field["id"]),
        "application_date",
    )
    inspection = get_latest_record(
        inspections_repository.list_inspections_by_field(field["id"]),
        "inspected_at",
    )
    spray_context = build_spray_context(spray_application)
    inspection_context = build_inspection_context(inspection)
    current_situation = "monitorar"

    return {
        "field_id": field["id"],
        "field_name": field["nome"],
        "generated_at": datetime.now(ZoneInfo(TIMEZONE)).isoformat(
            timespec="seconds"
        ),
        "current_situation": current_situation,
        "situation_label": SITUATION_LABELS[current_situation],
        "summary": "Situacao consolidada sem analise de risco disponivel.",
        "risk_context": {
            "main_disease": None,
            "risk_classification": None,
            "agronomic_index": None,
            "climate_index": None,
        },
        "spray_context": spray_context,
        "inspection_context": inspection_context,
        "recommended_next_action": SITUATION_ACTIONS[current_situation],
        "reasons": [
            "Nao foi possivel calcular a situacao atual completa para este talhao."
        ],
    }


def build_safe_field_situation(field: dict) -> dict:
    try:
        return build_current_field_situation(field)
    except Exception:
        return build_fallback_field_situation(field)


def is_recent_inspection(
    inspection_context: dict,
    current_date: date,
) -> bool:
    inspection_date = parse_optional_datetime(
        inspection_context.get("last_inspection_date")
        or inspection_context.get("inspected_at")
    )

    if inspection_date is None:
        return False

    return (current_date - inspection_date.date()).days <= RECENT_INSPECTION_WINDOW_DAYS


def build_distribution(
    counter: Counter,
    order: list[str],
    value_key: str,
) -> list[dict]:
    rows = [
        {
            value_key: key,
            "fields_count": counter[key],
        }
        for key in order
        if counter.get(key, 0) > 0
    ]
    remaining = sorted(
        key for key, count in counter.items() if key not in order and count > 0
    )

    rows.extend(
        {
            value_key: key,
            "fields_count": counter[key],
        }
        for key in remaining
    )

    return rows


def build_status_distribution(status_counter: Counter) -> list[dict]:
    return [
        {
            **row,
            "situation_label": SITUATION_LABELS.get(
                row["current_situation"],
                row["current_situation"],
            ),
        }
        for row in build_distribution(
            status_counter,
            SEASON_SITUATION_ORDER,
            "current_situation",
        )
    ]


def build_critical_field(situation: dict) -> dict:
    spray_context = situation["spray_context"]
    risk_context = situation["risk_context"]

    return {
        "field_id": situation["field_id"],
        "field_name": situation["field_name"],
        "current_situation": situation["current_situation"],
        "situation_label": situation["situation_label"],
        "main_disease": risk_context.get("main_disease"),
        "agronomic_index": risk_context.get("agronomic_index"),
        "estimated_defense_percent": spray_context.get(
            "estimated_defense_percent"
        ),
        "defense_status": spray_context.get("defense_status"),
        "recommended_next_action": situation["recommended_next_action"],
    }


def build_operational_alert(
    *,
    alert_type: str,
    severity: str,
    title: str,
    description: str,
    field_id: str | None = None,
    field_name: str | None = None,
) -> dict:
    alert = {
        "type": alert_type,
        "severity": severity,
        "title": title,
        "description": description,
    }

    if field_id is not None:
        alert["field_id"] = field_id

    if field_name is not None:
        alert["field_name"] = field_name

    return alert


def build_field_operational_alerts(
    situation: dict,
    current_date: date,
) -> list[dict]:
    alerts = []
    field_id = situation["field_id"]
    field_name = situation["field_name"]
    spray_context = situation["spray_context"]
    defense_status = spray_context.get("defense_status")

    if situation["current_situation"] == "prioridade_maxima":
        alerts.append(
            build_operational_alert(
                alert_type="prioridade_maxima",
                severity="alta",
                title="Talhao em prioridade maxima",
                description=(
                    "A situacao atual indica necessidade de verificacao operacional."
                ),
                field_id=field_id,
                field_name=field_name,
            )
        )

    if defense_status == "vencida":
        alerts.append(
            build_operational_alert(
                alert_type="defesa_vencida",
                severity="alta",
                title="Defesa estimada vencida",
                description=(
                    "A ultima pulverizacao registrada esta fora da janela operacional estimada."
                ),
                field_id=field_id,
                field_name=field_name,
            )
        )
    elif defense_status == "baixa":
        alerts.append(
            build_operational_alert(
                alert_type="defesa_baixa",
                severity="media",
                title="Defesa estimada baixa",
                description=(
                    "A defesa fitossanitaria estimada esta em nivel baixo para acompanhamento."
                ),
                field_id=field_id,
                field_name=field_name,
            )
        )
    elif defense_status == "sem_registro":
        alerts.append(
            build_operational_alert(
                alert_type="sem_pulverizacao",
                severity="media",
                title="Sem pulverizacao registrada",
                description=(
                    "Nao ha pulverizacao registrada para apoiar a leitura operacional do talhao."
                ),
                field_id=field_id,
                field_name=field_name,
            )
        )

    if not is_recent_inspection(situation["inspection_context"], current_date):
        alerts.append(
            build_operational_alert(
                alert_type="sem_inspecao_recente",
                severity="media",
                title="Sem inspecao recente",
                description=(
                    "Nao ha inspecao recente registrada para confirmar sintomas em campo."
                ),
                field_id=field_id,
                field_name=field_name,
            )
        )

    return alerts


def get_upcoming_calendar_events(
    current_date: date,
    days_ahead: int = UPCOMING_EVENT_WINDOW_DAYS,
) -> list[dict]:
    end_date = current_date + timedelta(days=days_ahead)
    upcoming_events = [
        event
        for event in calendar_events_repository.list_calendar_events()
        if current_date <= parse_calendar_event_date(event) <= end_date
    ]
    upcoming_events.sort(
        key=lambda event: (event["date"], event["end_date"], event["title"])
    )

    return upcoming_events[:10]


def build_calendar_operational_alerts(
    upcoming_events: list[dict],
    current_date: date,
) -> list[dict]:
    alerts = []

    for event in upcoming_events:
        if event.get("event_type") != "reaplicacao_prevista":
            continue

        event_date = parse_calendar_event_date(event)
        days_until_event = (event_date - current_date).days

        if days_until_event > 7:
            continue

        alerts.append(
            build_operational_alert(
                alert_type="reaplicacao_proxima",
                severity="alta" if days_until_event <= 1 else "media",
                title="Reaplicacao prevista proxima",
                description=(
                    f"{event['title']} prevista para {event['date']}."
                ),
                field_id=event.get("field_id"),
                field_name=event.get("field_name"),
            )
        )

    return alerts


@app.get("/safra/resumo")
def get_season_overview() -> dict:
    fields = fields_repository.list_fields()
    active_fields = [
        field for field in fields if field.get("status_lavoura") == "em_campo"
    ]
    current_date = datetime.now(ZoneInfo(TIMEZONE)).date()
    situations = [build_safe_field_situation(field) for field in active_fields]
    status_counter = Counter(
        situation["current_situation"] for situation in situations
    )
    defense_counter = Counter(
        situation["spray_context"].get("defense_status")
        for situation in situations
        if situation["spray_context"].get("defense_status")
    )
    estimated_defenses = [
        situation["spray_context"].get("estimated_defense_percent")
        for situation in situations
        if isinstance(
            situation["spray_context"].get("estimated_defense_percent"),
            (int, float),
        )
    ]
    operational_alerts = [
        alert
        for situation in situations
        for alert in build_field_operational_alerts(situation, current_date)
    ]
    upcoming_calendar_events = get_upcoming_calendar_events(current_date)
    operational_alerts.extend(
        build_calendar_operational_alerts(
            upcoming_calendar_events,
            current_date,
        )
    )
    season_metrics = get_season_metrics()
    average_estimated_defense_percent = (
        round(sum(estimated_defenses) / len(estimated_defenses), 2)
        if estimated_defenses
        else 0
    )
    critical_fields = [
        build_critical_field(situation)
        for situation in situations
        if situation["current_situation"] in CRITICAL_SITUATIONS
    ]
    critical_fields.sort(
        key=lambda field: (
            0 if field["current_situation"] == "prioridade_maxima" else 1,
            -(field.get("agronomic_index") or 0),
            field["field_name"],
        )
    )

    return {
        "summary": {
            "total_fields": len(fields),
            "active_fields": len(active_fields),
            "stable_fields": status_counter["estavel"],
            "monitoring_fields": (
                status_counter["monitorar"] + status_counter["monitorar_resposta"]
            ),
            "high_attention_fields": status_counter["alta_atencao"],
            "maximum_priority_fields": status_counter["prioridade_maxima"],
            "low_or_expired_defense_fields": len(
                [
                    situation
                    for situation in situations
                    if situation["spray_context"].get("defense_status")
                    in LOW_OR_EXPIRED_DEFENSE_STATUSES
                ]
            ),
            "fields_without_spray": len(
                [
                    situation
                    for situation in situations
                    if not situation["spray_context"].get("has_spray_record")
                ]
            ),
            "fields_without_recent_inspection": len(
                [
                    situation
                    for situation in situations
                    if not is_recent_inspection(
                        situation["inspection_context"],
                        current_date,
                    )
                ]
            ),
            "average_estimated_defense_percent": average_estimated_defense_percent,
        },
        "status_distribution": build_status_distribution(status_counter),
        "defense_distribution": build_distribution(
            defense_counter,
            SEASON_DEFENSE_ORDER,
            "defense_status",
        ),
        "critical_fields": critical_fields,
        "operational_alerts": operational_alerts,
        "top_product": (
            season_metrics["top_products"][0]
            if season_metrics["top_products"]
            else None
        ),
        "top_target": (
            season_metrics["top_targets"][0]
            if season_metrics["top_targets"]
            else None
        ),
        "upcoming_calendar_events": upcoming_calendar_events,
    }


def get_ranking_priority(
    risk_classification: str,
    management_relevance: str,
) -> str:
    if management_relevance != "ATIVA":
        return "SEM PRIORIDADE OPERACIONAL IMEDIATA"

    if risk_classification in ["CRÍTICO", "CRITICO", "CRÃTICO"]:
        return "PRIORIDADE MÁXIMA"
    if risk_classification == "ALTO":
        return "ALTA PRIORIDADE"
    if risk_classification == "MODERADO":
        return "MONITORAR"
    if risk_classification == "BAIXO":
        return "BAIXA PRIORIDADE"

    return "SEM PRIORIDADE OPERACIONAL IMEDIATA"


def calculate_historical_risk_bonus(field: dict) -> dict:
    bonus = 0
    reasons = []

    if field.get("historical_pressure") == "alta":
        bonus += 15
        reasons.append("Pressao historica alta da area")
    elif field.get("historical_pressure") == "media":
        bonus += 8
        reasons.append("Pressao historica media da area")

    if field.get("disease_incidence_level") == "alta":
        bonus += 15
        reasons.append("Incidencia historica alta")
    elif field.get("disease_incidence_level") == "media":
        bonus += 8
        reasons.append("Incidencia historica media")

    if field.get("had_disease_incidence") is True:
        bonus += 5
        reasons.append("Incidencia anterior de doenca registrada")

    peanut_repetition_years = field.get("peanut_repetition_years")
    if peanut_repetition_years is not None:
        if peanut_repetition_years >= 3:
            bonus += 10
            reasons.append(
                f"Repeticao de amendoim por {peanut_repetition_years} safras"
            )
        elif peanut_repetition_years == 2:
            bonus += 5
            reasons.append("Repeticao de amendoim por 2 safras")

    if field.get("crop_rotation") is False:
        bonus += 5
        reasons.append("Sem rotacao de cultura informada")

    return {
        "historical_risk_bonus": min(bonus, 30),
        "historical_risk_reasons": reasons,
    }


def build_ranking_item(
    field: dict,
    analysis: dict,
    spray_context: dict,
    inspection_context: dict,
) -> dict:
    risk = analysis["risk"]
    management_relevance = analysis["management_relevance"]["status"]
    historical_risk = calculate_historical_risk_bonus(field)
    agronomic_index_with_history = min(
        100,
        risk["agronomic_index"] + historical_risk["historical_risk_bonus"],
    )
    attention_priority = build_attention_priority(
        field=field,
        analysis=analysis,
        spray_context=spray_context,
        inspection_context=inspection_context,
    )

    return {
        "field_id": field["id"],
        "field_name": field["nome"],
        "city": field["cidade"],
        "crop_stage": analysis["field"]["crop_stage"],
        "disease": analysis["disease"],
        "pathogen": analysis["pathogen"],
        "risk_classification": risk["classification"],
        "climate_index": risk["climate_index"],
        "agronomic_index": risk["agronomic_index"],
        "agronomic_index_with_history": agronomic_index_with_history,
        **historical_risk,
        "management_relevance": management_relevance,
        "priority": (
            "SEM PRIORIDADE OPERACIONAL IMEDIATA"
            if management_relevance != "ATIVA"
            else PRIORITY_TO_LEGACY_PRIORITY[attention_priority["priority_label"]]
        ),
        **attention_priority,
        "main_action": analysis["actions"][0] if analysis["actions"] else None,
        "generated_at": analysis["generated_at"],
    }


def ranking_sort_key(item: dict) -> tuple:
    management_score = 1 if item["management_relevance"] == "ATIVA" else 0
    risk_class_score = RISK_CLASS_PRIORITY.get(item["risk_classification"], 0)

    return (
        management_score,
        item["priority_score"],
        item["confidence_score"],
        item["agronomic_index_with_history"],
        risk_class_score,
    )


@app.get("/ranking")
def get_ranking() -> dict:
    fields = fields_repository.list_fields()
    generated_at = datetime.now(ZoneInfo(TIMEZONE)).isoformat(timespec="seconds")
    ranking_items = []

    for field in fields:
        if not field.get("doencas_monitoradas"):
            continue

        analyses = run_compact_analyses_for_field(field)

        if not analyses:
            continue

        main_analysis = analyses[0]
        main_disease = main_analysis["disease"]
        spray_applications = (
            spray_applications_repository.list_spray_applications_by_field(
                field["id"]
            )
        )
        selected_spray_application = get_latest_matching_record(
            spray_applications,
            date_key="application_date",
            main_disease=main_disease,
            disease_key="target",
        )
        inspections = inspections_repository.list_inspections_by_field(field["id"])
        selected_inspection = get_latest_matching_record(
            inspections,
            date_key="inspected_at",
            main_disease=main_disease,
            disease_key="disease",
        )
        ranking_items.append(
            build_ranking_item(
                field,
                main_analysis,
                build_spray_context(selected_spray_application),
                build_inspection_context(selected_inspection),
            )
        )

    ranking_items.sort(key=ranking_sort_key, reverse=True)

    ranking = [
        {
            "rank": index,
            **item,
        }
        for index, item in enumerate(ranking_items, start=1)
    ]

    return {
        "generated_at": generated_at,
        "total_fields": len(fields),
        "total_items": len(ranking),
        "ranking": ranking,
    }


@app.post("/analisar")
def analyze_field(talhao: FieldAnalysisRequest) -> dict:
    talhao_dict = talhao.model_dump(mode="json")

    try:
        return risk_engine.run_disease_analysis_compact(talhao_dict)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    except HTTPException:
        raise
    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=f"Erro ao executar analise: {error}",
        ) from error
