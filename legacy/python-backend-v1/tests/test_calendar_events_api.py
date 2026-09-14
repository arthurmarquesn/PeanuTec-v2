from io import BytesIO
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api import app as api_app


client = TestClient(api_app.app)
TIMEZONE = ZoneInfo(api_app.TIMEZONE)
CALENDAR_EXPORT_HEADERS = [
    "Data",
    "Data final",
    "Talh\u00e3o",
    "Tipo de evento",
    "Status",
    "Produto",
    "Tipo do produto",
    "Alvo",
    "Dose",
    "Respons\u00e1vel",
    "Intervalo planejado",
    "Reaplica\u00e7\u00e3o prevista",
    "Observa\u00e7\u00f5es",
]


@pytest.fixture(autouse=True)
def isolated_files(tmp_path, monkeypatch):
    monkeypatch.setattr(
        api_app.fields_repository,
        "FIELDS_FILE",
        tmp_path / "outputs" / "fields.json",
    )
    monkeypatch.setattr(
        api_app.calendar_events_repository,
        "CALENDAR_EVENTS_FILE",
        tmp_path / "outputs" / "calendar_events.json",
    )
    monkeypatch.setattr(
        api_app.spray_applications_repository,
        "SPRAY_APPLICATIONS_FILE",
        tmp_path / "outputs" / "spray_applications.json",
    )
    monkeypatch.setattr(
        api_app.inspections_repository,
        "INSPECTIONS_FILE",
        tmp_path / "outputs" / "inspections.json",
    )


def make_field(
    field_id: str = "field-1",
    name: str = "Talhao A1",
    status_lavoura: str = "em_campo",
) -> dict:
    return {
        "id": field_id,
        "nome": name,
        "cidade": "Tupa-SP",
        "latitude": -21.9347,
        "longitude": -50.5136,
        "cultura": "Amendoim",
        "data_plantio": "2026-04-10",
        "status_lavoura": status_lavoura,
        "doencas_monitoradas": ["Mancha-preta"],
    }


def save_fields(fields: list[dict]) -> None:
    api_app.fields_repository.save_fields(fields)


def make_calendar_payload(**overrides) -> dict:
    payload = {
        "event_type": "monitoramento",
        "title": "Monitoramento pos-chuva",
        "field_id": "field-1",
        "date": "2026-06-17",
        "end_date": "2026-06-17",
        "product": "",
        "product_type": "outro",
        "target": "Mancha-preta",
        "planned_interval_days": 12,
        "notes": "Revisar folhas baixeiras.",
    }
    payload.update(overrides)
    return payload


def make_spray_payload(**overrides) -> dict:
    payload = {
        "application_date": (
            datetime.now(TIMEZONE) - timedelta(days=2)
        ).isoformat(timespec="seconds"),
        "product": "Clorotalonil",
        "target": "Mancha-preta",
        "dose": "1,5 L/ha",
        "responsible": "Tecnico responsavel",
        "planned_interval_days": 12,
        "notes": "Aplicacao preventiva realizada.",
    }
    payload.update(overrides)
    return payload


def make_inspection_payload(**overrides) -> dict:
    payload = {
        "disease": "Mancha-preta do amendoim",
        "symptoms_found": True,
        "visual_severity": "baixa",
        "defoliation_level": "baixa",
        "action_taken": "monitorar",
        "notes": "Inspecao geral da area.",
    }
    payload.update(overrides)
    return payload


def load_exported_calendar_rows(response) -> list[tuple]:
    workbook = load_workbook(BytesIO(response.content))
    worksheet = workbook.active

    return list(worksheet.iter_rows(values_only=True))


def test_post_calendario_eventos_creates_manual_event():
    field = make_field()
    save_fields([field])

    response = client.post("/calendario/eventos", json=make_calendar_payload())

    assert response.status_code == 200
    data = response.json()
    assert data["id"]
    assert data["event_type"] == "monitoramento"
    assert data["title"] == "Monitoramento pos-chuva"
    assert data["field_id"] == field["id"]
    assert data["field_name"] == field["nome"]
    assert data["date"] == "2026-06-17"
    assert data["end_date"] == "2026-06-17"
    assert data["color_key"] == "blue"
    assert data["status"] == "previsto"


def test_get_calendario_eventos_lists_events():
    save_fields([make_field()])
    created = client.post("/calendario/eventos", json=make_calendar_payload()).json()

    response = client.get("/calendario/eventos")

    assert response.status_code == 200
    data = response.json()
    assert data["total"] == 1
    assert data["events"] == [created]


def test_get_calendario_eventos_filters_by_field_id():
    save_fields(
        [
            make_field(field_id="field-1", name="Talhao A1"),
            make_field(field_id="field-2", name="Talhao B2"),
        ]
    )
    client.post("/calendario/eventos", json=make_calendar_payload(field_id="field-1"))
    client.post(
        "/calendario/eventos",
        json=make_calendar_payload(
            field_id="field-2",
            title="Monitoramento B2",
        ),
    )

    response = client.get("/calendario/eventos", params={"field_id": "field-2"})

    assert response.status_code == 200
    data = response.json()
    assert data["total"] == 1
    assert data["events"][0]["field_id"] == "field-2"


def test_get_calendario_eventos_filters_by_event_type():
    save_fields([make_field()])
    client.post("/calendario/eventos", json=make_calendar_payload())
    client.post(
        "/calendario/eventos",
        json=make_calendar_payload(
            event_type="observacao",
            title="Observacao geral",
            field_id=None,
        ),
    )

    response = client.get(
        "/calendario/eventos",
        params={"event_type": "observacao"},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["total"] == 1
    assert data["events"][0]["event_type"] == "observacao"


def test_get_calendario_exportar_excel_returns_excel_file():
    save_fields([make_field()])
    client.post(
        "/calendario/eventos",
        json=make_calendar_payload(
            event_type="pulverizacao",
            title="Pulverizacao A1",
            product="Fungicida X",
            product_type="fungicida",
        ),
    )

    response = client.get("/calendario/exportar-excel")

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "calendario-manejo-peanutec.xlsx" in response.headers[
        "content-disposition"
    ]
    rows = load_exported_calendar_rows(response)
    assert rows[0] == tuple(CALENDAR_EXPORT_HEADERS)
    assert rows[1][0] == "2026-06-17"
    assert rows[1][5] == "Fungicida X"


def test_get_calendario_exportar_excel_contains_expected_headers():
    response = client.get("/calendario/exportar-excel")

    assert response.status_code == 200
    rows = load_exported_calendar_rows(response)
    assert rows == [tuple(CALENDAR_EXPORT_HEADERS)]


def test_get_calendario_exportar_excel_filters_by_field_id():
    save_fields(
        [
            make_field(field_id="field-1", name="Talhao A1"),
            make_field(field_id="field-2", name="Talhao B2"),
        ]
    )
    client.post(
        "/calendario/eventos",
        json=make_calendar_payload(field_id="field-1", title="Evento A1"),
    )
    client.post(
        "/calendario/eventos",
        json=make_calendar_payload(field_id="field-2", title="Evento B2"),
    )

    response = client.get(
        "/calendario/exportar-excel",
        params={"field_id": "field-2"},
    )

    assert response.status_code == 200
    rows = load_exported_calendar_rows(response)
    assert len(rows) == 2
    assert rows[1][2] == "Talhao B2"


def test_get_calendario_exportar_excel_filters_by_event_type():
    save_fields([make_field()])
    client.post(
        "/calendario/eventos",
        json=make_calendar_payload(event_type="monitoramento"),
    )
    client.post(
        "/calendario/eventos",
        json=make_calendar_payload(
            event_type="observacao",
            title="Observacao geral",
            field_id=None,
        ),
    )

    response = client.get(
        "/calendario/exportar-excel",
        params={"event_type": "observacao"},
    )

    assert response.status_code == 200
    rows = load_exported_calendar_rows(response)
    assert len(rows) == 2
    assert rows[1][3] == "observacao"
    assert rows[1][2] is None


def test_get_calendario_exportar_excel_works_without_events():
    response = client.get("/calendario/exportar-excel")

    assert response.status_code == 200
    rows = load_exported_calendar_rows(response)
    assert rows == [tuple(CALENDAR_EXPORT_HEADERS)]


def test_get_calendario_resumo_returns_summary():
    save_fields(
        [
            make_field(field_id="field-1", status_lavoura="em_campo"),
            make_field(field_id="field-2", status_lavoura="colhido"),
        ]
    )
    future_date = (datetime.now(TIMEZONE).date() + timedelta(days=3)).isoformat()
    client.post(
        "/calendario/eventos",
        json=make_calendar_payload(
            event_type="pulverizacao",
            title="Pulverizacao realizada",
        ),
    )
    client.post(
        "/calendario/eventos",
        json=make_calendar_payload(
            event_type="inspecao",
            title="Inspecao realizada",
        ),
    )
    client.post(
        "/calendario/eventos",
        json=make_calendar_payload(
            event_type="reaplicacao_prevista",
            title="Reaplicacao prevista",
            date=future_date,
            end_date=future_date,
        ),
    )

    response = client.get("/calendario/resumo")

    assert response.status_code == 200
    assert response.json() == {
        "active_fields": 1,
        "events_count": 3,
        "spray_events_count": 1,
        "inspection_events_count": 1,
        "upcoming_attention_count": 1,
        "fields_without_recent_inspection": 0,
    }


def test_post_pulverizacoes_creates_spray_calendar_event():
    field = make_field()
    save_fields([field])

    spray_response = client.post(
        f"/talhoes/{field['id']}/pulverizacoes",
        json=make_spray_payload(),
    )
    events_response = client.get(
        "/calendario/eventos",
        params={"event_type": "pulverizacao"},
    )

    assert spray_response.status_code == 200
    assert events_response.status_code == 200
    events = events_response.json()["events"]
    assert len(events) == 1
    assert events[0]["event_type"] == "pulverizacao"
    assert events[0]["field_id"] == field["id"]
    assert events[0]["source_id"] == spray_response.json()["id"]
    assert events[0]["status"] == "realizado"
    assert events[0]["color_key"] == "yellow"


def test_post_pulverizacoes_with_generate_reapplication_false_does_not_create_reapplication_event():
    field = make_field()
    save_fields([field])

    spray_response = client.post(
        f"/talhoes/{field['id']}/pulverizacoes",
        json=make_spray_payload(generate_reapplication=False),
    )
    events_response = client.get(
        "/calendario/eventos",
        params={"event_type": "reaplicacao_prevista"},
    )

    assert spray_response.status_code == 200
    assert events_response.status_code == 200
    assert events_response.json()["events"] == []


def test_post_pulverizacoes_creates_reapplication_calendar_event():
    field = make_field()
    save_fields([field])
    application_date = datetime(2026, 6, 1, 8, 0, tzinfo=TIMEZONE)

    spray_response = client.post(
        f"/talhoes/{field['id']}/pulverizacoes",
        json=make_spray_payload(
            application_date=application_date.isoformat(timespec="seconds"),
            generate_reapplication=True,
            planned_interval_days=12,
        ),
    )
    events_response = client.get(
        "/calendario/eventos",
        params={"event_type": "reaplicacao_prevista"},
    )

    assert spray_response.status_code == 200
    events = events_response.json()["events"]
    assert len(events) == 1
    assert events[0]["date"] == "2026-06-13"
    assert events[0]["status"] == "previsto"
    assert events[0]["color_key"] == "purple"


def test_post_pulverizacoes_uses_explicit_reapplication_date():
    field = make_field()
    save_fields([field])
    application_date = datetime(2026, 6, 1, 8, 0, tzinfo=TIMEZONE)

    spray_response = client.post(
        f"/talhoes/{field['id']}/pulverizacoes",
        json=make_spray_payload(
            application_date=application_date.isoformat(timespec="seconds"),
            generate_reapplication=True,
            reapplication_date="2026-06-20",
            reapplication_notes="Retornar apos avaliacao do responsavel.",
        ),
    )
    events_response = client.get(
        "/calendario/eventos",
        params={"event_type": "reaplicacao_prevista"},
    )

    assert spray_response.status_code == 200
    events = events_response.json()["events"]
    assert len(events) == 1
    assert events[0]["date"] == "2026-06-20"
    assert events[0]["notes"] == "Retornar apos avaliacao do responsavel."


def test_post_pulverizacoes_uses_reapplication_interval_to_calculate_date():
    field = make_field()
    save_fields([field])
    application_date = datetime(2026, 6, 1, 8, 0, tzinfo=TIMEZONE)

    spray_response = client.post(
        f"/talhoes/{field['id']}/pulverizacoes",
        json=make_spray_payload(
            application_date=application_date.isoformat(timespec="seconds"),
            generate_reapplication=True,
            planned_interval_days=12,
            reapplication_interval_days=7,
        ),
    )
    events_response = client.get(
        "/calendario/eventos",
        params={"event_type": "reaplicacao_prevista"},
    )

    assert spray_response.status_code == 200
    events = events_response.json()["events"]
    assert len(events) == 1
    assert events[0]["date"] == "2026-06-08"
    assert events[0]["planned_interval_days"] == 7


def test_post_pulverizacoes_generate_reapplication_without_date_or_interval_returns_error():
    field = make_field()
    save_fields([field])
    payload = make_spray_payload(generate_reapplication=True)
    payload.pop("planned_interval_days")

    response = client.post(
        f"/talhoes/{field['id']}/pulverizacoes",
        json=payload,
    )

    assert response.status_code == 422
    assert "Informe a data prevista ou um intervalo" in str(response.json()["detail"])


def test_post_inspecoes_creates_inspection_calendar_event():
    field = make_field()
    save_fields([field])

    inspection_response = client.post(
        f"/talhoes/{field['id']}/inspecoes",
        json=make_inspection_payload(),
    )
    events_response = client.get(
        "/calendario/eventos",
        params={"event_type": "inspecao"},
    )

    assert inspection_response.status_code == 200
    events = events_response.json()["events"]
    assert len(events) == 1
    assert events[0]["event_type"] == "inspecao"
    assert events[0]["field_id"] == field["id"]
    assert events[0]["source_id"] == inspection_response.json()["id"]
    assert events[0]["status"] == "realizado"
    assert events[0]["color_key"] == "green"


def test_post_calendario_eventos_invalid_field_id_returns_error():
    response = client.post(
        "/calendario/eventos",
        json=make_calendar_payload(field_id="missing-field"),
    )

    assert response.status_code == 400
    assert "Talhao nao encontrado" in response.json()["detail"]
